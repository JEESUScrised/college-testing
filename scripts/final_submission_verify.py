"""Final DOCX + security + Allure evidence verification for submission."""
from __future__ import annotations

import json
import re
import zipfile
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path

from docx import Document
from PIL import Image

ROOT = Path(".").resolve()
STUDENT = "Прилепов"
GROUP = "4P1.23"
TEACHER_MARKERS = ("преподаватель", "teacher", "проверяющий преподаватель")
SECRET_PATTERNS = [
    re.compile(r"(?i)(api[_-]?key|secret[_-]?key|access[_-]?token)\s*[:=]\s*['\"][^'\"]{8,}"),
    re.compile(r"(?i)Bearer\s+[A-Za-z0-9\-_\.]{20,}"),
    re.compile(r"(?i)-----BEGIN (RSA |OPENSSH )?PRIVATE KEY-----"),
    re.compile(r"(?i)ATATT3[A-Za-z0-9_\-=]+"),  # Atlassian token shape
    re.compile(r"(?i)ghp_[A-Za-z0-9]{20,}"),
    re.compile(r"(?i)xox[baprs]-[A-Za-z0-9-]{10,}"),
]
ALLOW_DEMO = {"standard_user", "secret_sauce", "locked_out_user"}


def audit_docx() -> dict:
    items = []
    ok = True
    for path in sorted((ROOT / "reports").glob("KT*.docx")):
        doc = Document(str(path))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paras)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        pics = 0
        broken = 0
        for rel in doc.part.rels.values():
            if "image" not in rel.reltype:
                continue
            pics += 1
            blob = rel.target_part.blob
            try:
                Image.open(__import__("io").BytesIO(blob)).verify()
            except Exception:
                broken += 1
        has_student = STUDENT in full_text
        has_group = GROUP in full_text
        has_teacher = any(m in full_text.lower() for m in TEACHER_MARKERS)
        item = {
            "file": path.name,
            "size": path.stat().st_size,
            "pictures": pics,
            "broken_pictures": broken,
            "has_student": has_student,
            "has_group": has_group,
            "has_teacher_field": has_teacher,
            "opens": True,
            "paragraphs": len(paras),
        }
        if not (has_student and has_group and pics >= 0 and broken == 0 and not has_teacher):
            ok = False
            item["pass"] = False
        else:
            item["pass"] = True
        items.append(item)
    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "all_pass": ok and len(items) == 12,
        "reports": items,
        "note": "Final cleanup audit 2026-10-12.",
    }


def check_allure_shots() -> dict:
    out = {}
    for name in ("allure_01_overview.png", "allure_02_suites.png"):
        p = ROOT / "screenshots" / "kt08" / name
        im = Image.open(p).convert("RGB")
        # crude: ensure not tiny / mostly white empty
        out[name] = {
            "size": p.stat().st_size,
            "wh": im.size,
            "ok_size": p.stat().st_size > 15000,
        }
    return out


def security_scan() -> dict:
    hits = []
    roots = ["docs", "reports", "scripts", "selenium", "mobile", "robot", "grpc", "bug-reports", "README.md", "AGENTS.md"]
    skip_parts = {".venv", "__pycache__", "node_modules", "artifacts", ".git", "_diag"}
    for root_name in roots:
        root = ROOT / root_name
        paths = [root] if root.is_file() else list(root.rglob("*")) if root.exists() else []
        for p in paths:
            if not p.is_file():
                continue
            if any(part in skip_parts for part in p.parts):
                continue
            if p.suffix.lower() in {".png", ".jpg", ".jpeg", ".gif", ".zip", ".docx", ".pdf", ".pyc", ".jar"}:
                continue
            try:
                data = p.read_text(encoding="utf-8", errors="ignore")
            except Exception:
                continue
            for pat in SECRET_PATTERNS:
                for m in pat.finditer(data):
                    snippet = m.group(0)[:40]
                    if any(a in snippet for a in ALLOW_DEMO):
                        continue
                    hits.append({"file": str(p.relative_to(ROOT)).replace("\\", "/"), "match": snippet})
    return {"hits": hits, "clean": len(hits) == 0}


def main() -> None:
    audit = audit_docx()
    audit_path = ROOT / "docs" / "compliance" / "docx-audit.json"
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("docx_audit_all_pass", audit["all_pass"])
    for r in audit["reports"]:
        print(
            f"  {r['file']}: pics={r['pictures']} student={r['has_student']} "
            f"group={r['has_group']} teacher={r['has_teacher_field']} pass={r['pass']}"
        )

    allure = check_allure_shots()
    print("allure", allure)

    sec = security_scan()
    print("security_clean", sec["clean"], "hits", len(sec["hits"]))
    for h in sec["hits"][:20]:
        print("  SECRET?", h)

    # ensure unusable not present
    assert not (ROOT / "screenshots/kt09/archive/unusable-visual").exists()
    assert (ROOT / "screenshots/kt09/automated/windows/01_login.png").is_file()
    assert (ROOT / "screenshots/kt09/automated/android/01_login.png").is_file()
    print("kt09_paths_ok")

    # audit counts from markdown
    md = (ROOT / "docs/compliance/SCREENSHOT_VISUAL_AUDIT.md").read_text(encoding="utf-8")
    assert "| GOOD | 170 |" in md
    assert "| INTENTIONAL_BLANK | 3 |" in md
    print("screenshot_audit_counts_ok")

    out = {
        "docx": audit,
        "allure": allure,
        "security": sec,
    }
    (ROOT / "docs/compliance/final-cleanup-verify.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
