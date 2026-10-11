"""Generate reports/KT04.docx from verified KT04 + real Jira evidence."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot, add_table
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt04"
JIRA = SHOTS / "jira"
JIRA_FINAL = JIRA / "final"
OUTPUT = Path(__file__).resolve().parent / "KT04.docx"

SCREENSHOTS = [
    ("tc01_valid_login.png", "Рисунок 1 — TC-01 PASS: валидный вход"),
    ("tc02_seed001_empty_password.png", "Рисунок 2 — SEED-001: пустой пароль принят"),
    ("tc03_seed002_invalid_email.png", "Рисунок 3 — SEED-002: email user@mail принят"),
    ("tc04_seed003_zero_quantity.png", "Рисунок 4 — SEED-003: qty=0 принят"),
    ("tc05_valid_order.png", "Рисунок 5 — TC-05 PASS: валидный заказ"),
]

LIFECYCLE = SHOTS / "lifecycle"
LIFE_SHOTS = [
    (LIFECYCLE / "retest_seed001_empty_password.png", "Рисунок 6 — Lifecycle retest SEED-001 PASS"),
    (LIFECYCLE / "retest_seed002_invalid_email.png", "Рисунок 7 — Lifecycle retest SEED-002 PASS"),
    (LIFECYCLE / "retest_seed003_zero_quantity.png", "Рисунок 8 — Lifecycle retest SEED-003 PASS"),
]

JIRA_SHOTS = [
    (JIRA / "01_board_college_qa.png", "Рисунок 9 — Jira board College QA Defect Lifecycle (KAN)"),
    (JIRA_FINAL / "02_KAN-4_issue_comments.png", "Рисунок 10 — KAN-4 Done + 3 comments (remote verify)"),
    (JIRA_FINAL / "07_saved_filter_bugs_done.png", "Рисунок 11 — Saved filter KT04 KAN Bugs Done"),
    (JIRA_FINAL / "04_project_summary.png", "Рисунок 12 — Project Summary: Выполнено 3"),
    (JIRA_FINAL / "08_dashboard_configured_stats.png", "Рисунок 13 — Dashboard: status/priority charts (configured)"),
]


def build() -> Path:
    doc = Document()
    setup_a4(doc)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №4. Составление отчетов о дефектах и работа с баг-трекером Jira",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке", center=True)
    add_para(doc, "Дата обновления: 2026-10-11 (remote Jira verification)", center=True)

    add_heading(doc, "Что нужно было закрыть")
    add_para(
        doc,
        "КТ про баг-трекер: завести проект в Jira, настроить типы/приоритеты/workflow, "
        "оформить баг-репорты, назначить, провести по стадиям, «починить», сделать retest "
        "и закрыть. В конце — фильтр, summary и dashboard со статистикой. Я делал это на "
        "реальном облачном сайте, не на скриншотах из учебника.",
    )

    add_heading(doc, "Мой проект в Jira")
    add_para(doc, "Site: https://collegeqakt04.atlassian.net")
    add_para(doc, "Project: College QA Defect Lifecycle, key KAN (team-managed / next-gen).")
    add_para(doc, "Board: https://collegeqakt04.atlassian.net/jira/software/projects/KAN/boards/2")
    add_para(doc, "Issue types: Epic, Story, Task, Баг, Subtask.")
    add_para(doc, "Workflow: Open → In Progress → Ready for Testing → Done.")
    add_para(
        doc,
        "С Components вышла засада: на team-managed проекте KAN они UNAVAILABLE. "
        "POST /components отдаёт 405, в editmeta поля components нет. Вместо Components "
        "поставил labels: authentication / registration / order-processing — это замена, "
        "а не «настоящие» Components. Проект ради Components на company-managed не "
        "мигрировал: ключ KAN уже заведён, баги на нём.",
    )

    add_heading(doc, "Три бага: KAN-4, KAN-5, KAN-6")
    add_para(
        doc,
        "Все три заведены как тип Баг, прогнаны по workflow до Done, у каждого по три "
        "комментария (triage / fix / retest). Проверял через REST.",
    )
    add_para(doc, "KAN-4 — Login accepts empty password — High — Done — assignee JEESUS — 3 comments")
    add_para(doc, "https://collegeqakt04.atlassian.net/browse/KAN-4")
    add_para(doc, "KAN-5 — Registration accepts email without TLD — High — Done — 3 comments")
    add_para(doc, "https://collegeqakt04.atlassian.net/browse/KAN-5")
    add_para(doc, "KAN-6 — Order accepts quantity zero — Medium — Done — 3 comments")
    add_para(doc, "https://collegeqakt04.atlassian.net/browse/KAN-6")
    add_para(doc, "Changelog KAN-4: Open→In Progress→Ready for Testing→Done.")
    add_para(doc, "GitHub #1–#3 сохранены как исторические записи.")

    add_heading(doc, "Seeded suite до фиксов (исторический)")
    add_para(
        doc,
        "Campus Portal на fixtures/kt04 специально «дырявый»: два валидных сценария "
        "проходят, три seed-бага падают — так и задумано для учебного FAIL.",
    )
    add_para(doc, "Campus Portal fixtures/kt04: 2 PASS / 3 FAIL (ожидаемые учебные FAIL).")
    add_code(doc, "======================== 3 failed, 2 passed in 18.08s =========================")

    add_heading(doc, "Скриншоты seeded suite")
    for name, caption in SCREENSHOTS:
        add_shot(doc, SHOTS / name, caption)

    add_heading(doc, "Lifecycle-фиксы и Selenium retest")
    add_para(
        doc,
        "В fixtures/kt04_lifecycle закрыл дыры: пустой пароль больше не принимается, "
        "email без TLD режется, qty должен быть > 0. После этого retest зелёный.",
    )
    add_para(doc, "fixtures/kt04_lifecycle: empty password / email TLD / qty>0. Retest 3 passed in 8.74s.")
    add_code(doc, "============================== 3 passed in 8.74s ==============================")
    for path, caption in LIFE_SHOTS:
        add_shot(doc, path, caption)

    add_heading(doc, "Назначения, комментарии, переходы")
    add_para(
        doc,
        "На каждом баге: assign JEESUS; 3 комментария (triage / fix / retest); "
        "переходы Open→In Progress→Ready for Testing→Done. "
        "Все роли на одном аккаунте JEESUS Cursor-агентом (образовательная симуляция) — "
        "отдельных «разработчика» и «тестировщика» в учебном проекте нет.",
    )

    add_heading(doc, "Скриншоты из Jira")
    add_para(
        doc,
        "Доска, карточка KAN-4 с комментариями, saved filter, project summary и "
        "dashboard с графиками — всё с remote verification 2026-10-11.",
    )
    for path, caption in JIRA_SHOTS:
        add_shot(doc, path, caption)

    add_heading(doc, "Фильтры, summary, dashboard")
    add_para(doc, "Saved filter: https://collegeqakt04.atlassian.net/issues/?filter=10002")
    add_para(doc, "Dashboard: https://collegeqakt04.atlassian.net/jira/dashboards/10001")
    add_para(doc, "Summary: https://collegeqakt04.atlassian.net/jira/software/projects/KAN/summary")
    add_para(doc, "Totals (re-verified): Bugs=3, Open=0, Done=3, High=2, Medium=1.")
    add_para(
        doc,
        "Dashboard gadgets configured on filter 10002: Bugs by Status (ГОТОВО=3), "
        "Bugs by Priority (High=2 / Medium=1), Priority×Status, KAN Bugs list.",
    )
    add_para(doc, "Limitation: native Components unavailable (405); Components breakdown N/A.")

    add_heading(doc, "Сверка с требованиями")
    add_table(
        doc,
        ["Требование", "Статус", "Доказательство"],
        [
            ("Ознакомление с Jira", "PASS", "collegeqakt04 + REST verify"),
            ("Создание проекта", "PASS", "KAN College QA Defect Lifecycle"),
            ("Типы/приоритеты/workflow", "PASS", "Баг + Ready for Testing"),
            ("Детальные баг-репорты", "PASS", "KAN-4/5/6 descriptions"),
            ("Назначения / комментарии", "PASS", "JEESUS + 3 comments/issue"),
            ("Переходы workflow", "PASS", "changelog → Done"),
            ("Исправление + retest", "PASS", "kt04_lifecycle 3/3"),
            ("Закрытие в Jira", "PASS", "Готово"),
            ("Native Components", "PARTIAL", "UNAVAILABLE on team-managed; labels substitute"),
            ("Статистика / reports", "PASS", "filter + configured dashboard charts"),
        ],
    )

    add_heading(doc, "Заключение")
    add_para(
        doc,
        "В Jira на collegeqakt04.atlassian.net лежат три бага KAN-4 / KAN-5 / KAN-6 — "
        "все Done, у каждого assignee, три комментария и полный проход "
        "Open→In Progress→Ready for Testing→Done. Локально после фиксов в "
        "kt04_lifecycle Selenium retest дал 3 passed in 8.74s. Исторический seeded прогон "
        "3 failed, 2 passed и GitHub #1–#3 не трогал — это «до» исправлений. "
        "Единственное честное ограничение: Components на team-managed UNAVAILABLE, "
        "поэтому метки labels вместо Components.",
    )

    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    required = [
        "работа с баг-трекером Jira",
        "collegeqakt04.atlassian.net",
        "KAN-4",
        "KAN-5",
        "KAN-6",
        "3 passed in 8.74s",
        "3 failed, 2 passed",
        "UNAVAILABLE",
        "Заключение",
    ]
    missing = [r for r in required if r not in text]
    if missing:
        raise AssertionError(missing)
    pictures = list(doc.inline_shapes)
    if len(pictures) < 10:
        raise AssertionError(f"pictures={len(pictures)}")
    images = [r for r in doc.part.rels.values() if r.reltype == RT.IMAGE]
    if not images:
        raise AssertionError("no images")
    print(f"VERIFY_OK size={path.stat().st_size} pictures={len(pictures)} unique={len(images)}")


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    verify(out)
