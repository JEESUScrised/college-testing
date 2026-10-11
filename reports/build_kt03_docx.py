"""Generate reports/KT03.docx — historical local PO + public Sauce PO remediation."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt03"
PUBLIC = SHOTS / "public"
OUTPUT = Path(__file__).resolve().parent / "KT03.docx"

HIST_SHOTS = [
    ("01_po_main_before_open.png", "Рисунок 1 — Главное окно (локальный PO)"),
    ("03_po_switched_to_secondary.png", "Рисунок 2 — Вторичное окно через WindowsSecondaryPage"),
    ("04_po_secondary_content_verified.png", "Рисунок 3 — Проверка содержимого вторичного окна"),
    ("06_po_returned_to_main.png", "Рисунок 4 — Возврат в исходное окно"),
    ("07_po_iframe_host.png", "Рисунок 5 — Host с iframe (IframePage)"),
    ("08_po_iframe_interaction.png", "Рисунок 6 — Взаимодействие внутри iframe"),
    ("09_po_back_to_default_content.png", "Рисунок 7 — leave_iframe / default_content"),
]

PUBLIC_SHOTS = [
    ("po_01_login_form.png", "Рисунок 8 — SauceLoginPage: форма login"),
    ("po_02_inventory_via_po.png", "Рисунок 9 — SauceInventoryPage после login"),
    ("po_03_product_details.png", "Рисунок 10 — SauceProductPage"),
    ("po_04_cart.png", "Рисунок 11 — SauceCartPage"),
    ("po_05_checkout_complete.png", "Рисунок 12 — SauceCheckoutPage complete"),
    ("po_06_logout.png", "Рисунок 13 — logout через Page Object"),
]

KT03_HIST = """\
============================= 4 passed in 14.90s ==============================
selenium/tests/test_kt03_page_object.py (windows/iframe Page Objects)
"""

KT03_PUBLIC = """\
============================= test session starts =============================
platform win32 -- Python 3.13.2, pytest-8.4.2
collecting ... collected 5 items
selenium/tests/test_kt03_sauce_po.py::test_po_login_page_opens_and_authenticates PASSED
selenium/tests/test_kt03_sauce_po.py::test_po_inventory_sort_and_open_product PASSED
selenium/tests/test_kt03_sauce_po.py::test_po_product_add_and_cart PASSED
selenium/tests/test_kt03_sauce_po.py::test_po_checkout_flow PASSED
selenium/tests/test_kt03_sauce_po.py::test_po_menu_logout PASSED
============================= 5 passed in 21.11s ==============================
"""


def build() -> Path:
    missing = [n for n, _ in HIST_SHOTS if not (SHOTS / n).is_file()]
    missing += [n for n, _ in PUBLIC_SHOTS if not (PUBLIC / n).is_file()]
    if missing:
        raise FileNotFoundError(missing)

    doc = Document()
    setup_a4(doc)

    add_title_fields(doc, add_para)
    add_para(doc, "КТ №3. Паттерн Page Object", bold=True, center=True)
    add_para(doc, "Отчёт (включая remediation 2026-10-10)", center=True)

    add_heading(doc, "Зачем мне Page Object")
    add_para(
        doc,
        "В КТ №2 локаторы уже начали дублироваться между кейсами. Page Object как раз "
        "про это: страница = класс, кнопки и поля — поля/методы класса, а в тесте остаётся "
        "сценарий и assertions. Если на Sauce Demo сменится id кнопки «Add to cart», правка "
        "одна — в SauceInventoryPage, а не в пяти тестах сразу.",
    )
    add_para(
        doc,
        "По заданию нужно было применить паттерн. Я сделал публичный PO на "
        "https://www.saucedemo.com/ и оставил локальный PO для окон/iframe — тот же подход, "
        "другие страницы.",
    )

    add_heading(doc, "Как разложил классы")
    add_para(
        doc,
        "Общий BasePage держит driver, ожидания и мелкие хелперы. Дальше — конкретные "
        "страницы. Тесты только вызывают методы вроде login_as(...) или add_first_product().",
    )
    add_para(doc, "• Исторический: BasePage, WindowsMainPage/WindowsSecondaryPage, IframePage")
    add_para(doc, "• Remediation: selenium/pages/sauce/* (Login, Inventory, Product, Cart, Checkout)")
    add_para(doc, "• Тесты: test_kt03_page_object.py + test_kt03_sauce_po.py")

    add_heading(doc, "Что прогонял")
    add_para(
        doc,
        "На Sauce Demo цепочка через PO: login → inventory (сортировка, открытие товара) → "
        "корзина → checkout → logout из меню. На локальных фикстурах — те же окна и iframe, "
        "что в КТ 02, только через page-классы.",
    )
    add_code(
        doc,
        """# Публичный Sauce Demo PO
.\\.venv\\Scripts\\python.exe -m pytest selenium\\tests\\test_kt03_sauce_po.py -v --browser=chrome --page-load-strategy=eager

# Исторический локальный PO
.\\.venv\\Scripts\\python.exe -m pytest selenium\\tests\\test_kt03_page_object.py -v --browser=chrome""",
    )

    add_heading(doc, "Результаты прогонов")
    add_para(doc, "Публичный PO (2026-10-10):")
    add_code(doc, KT03_PUBLIC)
    add_para(doc, "Исторический PO (сохранён):")
    add_code(doc, KT03_HIST)
    add_para(
        doc,
        "Публичный слой: 5 passed in 21.11s. Локальный PO окон/iframe: 4 passed in 14.90s.",
    )

    add_heading(doc, "Доказательства — локальный PO")
    for name, caption in HIST_SHOTS:
        add_shot(doc, SHOTS / name, caption)

    add_heading(doc, "Доказательства — Sauce Demo через PO")
    for name, caption in PUBLIC_SHOTS:
        add_shot(doc, PUBLIC / name, caption)

    add_heading(doc, "Заключение")
    add_para(
        doc,
        "Page Object оправдал себя на живом saucedemo: пять сценариев от login до logout "
        "идут через page-классы, прогон 5 passed in 21.11s. Локальный PO для окон/iframe "
        "тоже зелёный — 4 passed in 14.90s. Главный плюс для меня: тест читается как "
        "пользовательский сценарий, а селекторы не размазаны по файлам. Дальше те же "
        "sauce-страницы удобно переиспользовать в следующих КТ.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    for item in ("КТ №3. Паттерн Page Object", "5 passed in 21.11s", "4 passed in 14.90s", "saucedemo", "Заключение"):
        if item.lower() not in text.lower() and item not in text:
            # saucedemo may appear only in code blocks as URL path — check soft
            if item == "saucedemo" and "Sauce" in text:
                continue
            if item not in text:
                raise AssertionError(f"missing {item}")
    expected = len(HIST_SHOTS) + len(PUBLIC_SHOTS)
    if len(list(doc.inline_shapes)) < expected:
        raise AssertionError("not enough pictures")
    if not any(rel.reltype == RT.IMAGE for rel in doc.part.rels.values()):
        raise AssertionError("no images")
    print(f"VERIFY_OK size={path.stat().st_size}")


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    verify(out)
