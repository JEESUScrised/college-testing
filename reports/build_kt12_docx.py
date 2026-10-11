"""Generate reports/KT12.docx — gRPC InventoryService pytest suite."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt12"
OUTPUT = Path(__file__).resolve().parent / "KT12.docx"
PROTO = ROOT / "grpc" / "proto" / "inventory.proto"

RESULTS = [
    ("TC-12-01", "CreateItem — успех", "PASS"),
    ("TC-12-02", "GetItem — существующий", "PASS"),
    ("TC-12-03", "дубликат → ALREADY_EXISTS", "PASS"),
    ("TC-12-04", "невалидные данные → INVALID_ARGUMENT", "PASS"),
    ("TC-12-05", "нет записи → NOT_FOUND", "PASS"),
    ("TC-12-06", "UpdateStock + повторный Get", "PASS"),
    ("TC-12-07", "ListItems — server stream", "PASS"),
    ("TC-12-08", "BulkCreate — client stream", "PASS"),
    ("TC-12-09", "EchoWatch — bidi stream", "PASS"),
    ("TC-12-10", "SlowPing → DEADLINE_EXCEEDED", "PASS"),
]


def build() -> Path:
    doc = Document()
    setup_a4(doc, left=2.5, right=1.5)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №12. Автоматизированное тестирование реального gRPC-сервиса",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке (финальная)", center=True)
    add_para(doc, "Дата: 2026-10-10", center=True)

    add_heading(doc, "1. Цель")
    add_para(
        doc,
        "Собрать локальный InventoryService на Python + Protocol Buffers, "
        "сгенерировать stubs через grpc_tools.protoc и написать ровно десять "
        "pytest-тестов, которые реально ходят по HTTP/2 — без моков канала.",
    )

    add_heading(doc, "2. gRPC и proto — своими словами")
    add_para(
        doc,
        "gRPC — это RPC поверх HTTP/2. Контракт лежит в .proto (у меня proto3): "
        "сообщения и сервис. protoc потом генерит сериализацию и stubs для клиента "
        "и сервера. Удобно, что типы и методы описаны в одном месте.",
    )

    add_heading(doc, "3. Как клиент ходит на сервер")
    add_para(
        doc,
        "В тестах поднимаю insecure_channel на 127.0.0.1 и свободный порт, жду READY "
        "и дергаю InventoryServiceStub. По проводам — HTTP/2, тело — protobuf.",
    )
    add_para(
        doc,
        "Учебный сервис без TLS. Для продакшена так нельзя — нужны шифрование и auth.",
        bold=True,
    )

    add_heading(doc, "4. Архитектура InventoryService")
    add_para(
        doc,
        "Цепочка простая: proto → generated stubs → InventoryServicer + "
        "InventoryStore (in-memory, thread-safe) → pytest-фикстура с lifecycle "
        "сервера → десять тестов. Стор специально в памяти, чтобы каждый прогон "
        "был изолированным.",
    )

    add_heading(doc, "5. Контракт .proto и генерация")
    add_para(doc, "Фрагмент service InventoryService из inventory.proto:")
    proto_text = PROTO.read_text(encoding="utf-8") if PROTO.is_file() else ""
    start = proto_text.find("service InventoryService")
    end = proto_text.find("}", start) + 1 if start >= 0 else 0
    add_code(doc, proto_text[start:end] if start >= 0 else "(proto missing)")
    add_para(
        doc,
        "Генерил скриптом scripts/generate_kt12_proto.ps1 → "
        "grpc/generated/inventory_pb2.py и inventory_pb2_grpc.py "
        "(пометки DO NOT EDIT не трогал).",
    )

    add_heading(doc, "6. Сервер и клиент")
    add_para(
        doc,
        "server/service.py — валидация и context.abort(...) с нужными status codes. "
        "client/cli.py — маленький демо-вызов CreateItem + GetItem. "
        "Сервер руками: python -m server --port 50051 (из каталога с нормальным PYTHONPATH).",
    )

    add_heading(doc, "7. Unary RPC")
    add_para(
        doc,
        "CreateItem, GetItem, UpdateStock, SlowPing — классика «один запрос → один ответ». "
        "Поля и коды ответов дублировал в evidence.jsonl, чтобы потом не гадать по логу.",
    )

    add_heading(doc, "8. Streaming")
    add_para(
        doc,
        "ListItems — server streaming: сервер отдаёт итератор Item. "
        "BulkCreate — client streaming: клиент шлёт пачку, в ответ BulkCreateResponse. "
        "EchoWatch — bidi: на каждый WatchRequest ровно один WatchResponse, порядок тот же.",
    )

    add_heading(doc, "9. Ошибки и status codes")
    add_para(
        doc,
        "ALREADY_EXISTS — если id уже есть, INVALID_ARGUMENT — пустое имя или quantity < 0, "
        "NOT_FOUND — запросили то, чего нет. Это как раз то, что хочется видеть в API-тестах.",
    )

    add_heading(doc, "10. Deadline и отмена")
    add_para(
        doc,
        "Для TC-12-10 вызываю SlowPing(delay_ms=800) с client timeout=0.1s — "
        "ожидаемо ловлю DEADLINE_EXCEEDED. На сервере крутится проверка "
        "context.is_active() раз в 20 мс, чтобы не висеть зря после отмены.",
    )

    add_heading(doc, "11. Таблица тест-кейсов")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "ID", "Сценарий", "Статус"
    for row in RESULTS:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    add_para(doc, "")
    add_para(
        doc,
        "Входы и ожидаемые коды расписаны в docs/kt12/test-cases.md. "
        "Фикстура inventory_env каждый раз поднимает свой сервер на свободном порту.",
    )

    add_heading(doc, "12. Результат pytest")
    add_para(doc, "Сначала отдельно проверил TC-12-01 — PASS. Потом весь suite:")
    add_code(doc, "============================= 10 passed in 2.34s ==============================")
    add_para(doc, "JUnit: tests=10, failures=0, errors=0, skipped=0.")

    add_heading(doc, "13. Артефакты")
    add_para(
        doc,
        "В grpc/artifacts/kt12/: лог pytest, junit XML, HTML-отчёт, evidence.jsonl, "
        "versions JSON. Контрольные SHA16: log=E656BF1468AA2C30, "
        "junit=33F69FACE4CA7013, html=7329D8123EB2800D.",
    )
    add_code(
        doc,
        r""".\scripts\run_kt12.ps1
# evidence excerpt:
# DEADLINE_EXCEEDED SlowPing timeout=0.1 delay_ms=800
# ListItems ids=["s1","s2","s3"]
# EchoWatch FOUND/MISSING/FOUND""",
    )
    add_shot(
        doc,
        SHOTS / "01_pytest_html_report.png",
        "Рисунок 1 — pytest-html отчёт KT12 (10 passed)",
        width_cm=14.0,
    )

    add_heading(doc, "14. Заключение")
    add_para(
        doc,
        "КТ 12 сделана: живой gRPC InventoryService, generated stubs, unary + "
        "server/client/bidi streaming, проверка DEADLINE_EXCEEDED. Suite — "
        "10 passed in 2.34s. Ограничения понятные: insecure channel, in-memory store, "
        "без облака и Docker. Это последняя контрольная точка курса (КТ01–КТ12).",
    )

    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    required = [
        "КТ №12. Автоматизированное тестирование реального gRPC-сервиса",
        "10 passed in 2.34s",
        "DEADLINE_EXCEEDED",
        "Заключение",
    ]
    missing = [r for r in required if r not in text]
    if missing:
        raise SystemExit(f"DOCX missing: {missing}")
    print(f"OK {path} size={path.stat().st_size} chars={len(text)}")


if __name__ == "__main__":
    out = build()
    verify(out)
