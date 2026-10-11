"""Generate reports/KT09.docx — SauceDemo Windows/Android + Appium gestures."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot, add_table
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt09"
AUTO_WIN = SHOTS / "automated" / "windows"
AUTO_AND = SHOTS / "automated" / "android"
OUTPUT = Path(__file__).resolve().parent / "KT09.docx"

GESTURE_RESULTS = [
    ("TC-09-01", "свайп вверх по Views", "PASS"),
    ("TC-09-02", "свайп вниз по Views", "PASS"),
    ("TC-09-03", "scroll до TextFields и открытие", "PASS"),
    ("TC-09-04", "горизонтальный свайп влево (Gallery)", "PASS"),
    ("TC-09-05", "горизонтальный свайп вправо (Gallery)", "PASS"),
]

COMPARE = [
    (
        "Страница логина",
        "01_login.png",
        "Форма Swag Labs на месте: Username, Password, кнопка Login.",
        "PASS",
        "PASS",
        "Одинаковый публичный UI SauceDemo, расхождений по смыслу нет.",
    ),
    (
        "Каталог / Products",
        "02_inventory.png",
        "Список товаров с названиями и ценами.",
        "PASS",
        "PASS",
        "На десктопе сетка, на Android — узкая одноколоночная вёрстка.",
    ),
    (
        "Карточка товара",
        "03_product_details.png",
        "Название, описание, цена, Add to cart / Remove.",
        "PASS",
        "PASS",
        "Тот же DOM inventory-item / inventory-details.",
    ),
    (
        "Товар в корзине",
        "04_added_to_cart.png",
        "После Add to cart бейдж корзины = 1.",
        "PASS",
        "PASS",
        "Поведение совпадает.",
    ),
    (
        "Содержимое корзины",
        "05_cart.png",
        "Your Cart: товар, количество, цена.",
        "PASS",
        "PASS",
        "Поведение совпадает.",
    ),
    (
        "Боковое меню",
        "06_menu.png",
        "Burger menu: All Items, About, Logout.",
        "PASS",
        "PASS",
        "На Android меню важнее — экран узкий, пункты спрятаны.",
    ),
    (
        "Форма checkout",
        "07_checkout_form.png",
        "Checkout Information: First Name / Last Name / Zip.",
        "PASS",
        "PASS",
        "Совпадает; на Android suite ещё ловится валидация пустых полей.",
    ),
    (
        "Подтверждение заказа",
        "08_order_confirmation.png",
        "Checkout Complete — Thank you for your order.",
        "PASS",
        "PASS",
        "Ок; на части Android-кадров бывают артефакты GPU эмулятора.",
    ),
]


def build() -> Path:
    doc = Document()
    setup_a4(doc, left=2.0, right=1.5)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №9. Кроссплатформенное тестирование (SauceDemo) и жесты Appium",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке", center=True)
    add_para(doc, "Дата: 2026-10-11", center=True)

    add_heading(doc, "1. Зачем делал эту КТ")
    add_para(
        doc,
        "Нужно было прогнать публичный SauceDemo и на Windows, и на Android, "
        "снять автоматические скриншоты по одним и тем же шагам и отдельно "
        "пощупать жесты Appium на ApiDemos. Ручной проход человеком я в этот "
        "раз не закрывал — ниже это честно отмечено как PENDING.",
    )

    add_heading(doc, "2. Что тестировал и на чём")
    add_para(doc, "Сайт: https://www.saucedemo.com/ (Swag Labs).")
    add_para(
        doc,
        "Windows 10, Mozilla Firefox через Selenium 4, окно 1280×900. "
        "Кадры лежат в screenshots/kt09/automated/windows/ — всего 8 штук.",
    )
    add_para(
        doc,
        "Android: эмулятор emulator-5554 (Android 16 / API 36), Chrome 134.0.6998.135, "
        "ChromeDriver 134.0.6998.165, Appium 3.8. Скриншоты — "
        "screenshots/kt09/automated/android/ (тоже 8). Их снимал DOM-suite TC09-W* "
        "плюс отдельный захват checkout-complete.",
    )
    add_para(
        doc,
        "Важно: все скриншоты в отчёте — AUTOMATED evidence. "
        "Manual verification: PENDING HUMAN VERIFICATION.",
        bold=True,
    )

    add_heading(doc, "3. Что уже проходило раньше (цифры сохранены)")
    add_para(
        doc,
        "Жесты на ApiDemos: 5 passed in 54.93s. "
        "Старый Windows cross-platform web suite: 2 passed in 12.43s. "
        "Android Chrome DOM по SauceDemo: 11 passed in 36.91s "
        "(TC09-W01…W10 и journey).",
    )
    add_code(
        doc,
        "Windows Firefox visual capture: 8 screenshots (2026-10-11)\n"
        "Android Chrome automated visual set: 8 screenshots "
        "(Appium + ChromeDriver web screenshot, emulator soft GPU; "
        "scripts/capture_kt09_android_visual.py)\n"
        "Manual checklist: PENDING HUMAN VERIFICATION",
    )

    add_heading(doc, "4. Сравнение Windows ↔ Android по сценариям")
    add_para(
        doc,
        "Ниже восемь одинаковых шагов. Для каждого — что ждал, что вышло "
        "по автодоказательствам, скрин Windows Firefox и скрин Android Chrome. "
        "Подписи помечены [AUTOMATED], чтобы не перепутать с ручными.",
    )

    for i, (title, fname, expected, win_r, and_r, conclusion) in enumerate(COMPARE, start=1):
        add_heading(doc, f"4.{i}. {title}", level=2)
        add_para(doc, f"Ожидал: {expected}")
        add_para(
            doc,
            f"Windows Firefox — {win_r}. Android Chrome — {and_r}. "
            f"Вывод: {conclusion}",
        )
        add_para(doc, "Windows (automated):", bold=True)
        add_shot(
            doc,
            AUTO_WIN / fname,
            f"Рис. W-{i} — Windows Firefox — {title} [AUTOMATED]",
            width_cm=14.0,
        )
        add_para(doc, "Android (automated):", bold=True)
        add_shot(
            doc,
            AUTO_AND / fname,
            f"Рис. A-{i} — Android Chrome — {title} [AUTOMATED]",
            width_cm=8.5,
        )

    add_heading(doc, "5. Сводная таблица")
    add_table(
        doc,
        ["Сценарий", "Windows", "Android", "Файл", "Вывод"],
        [
            (title, win_r, and_r, fname, conclusion)
            for title, fname, _exp, win_r, and_r, conclusion in COMPARE
        ],
    )

    add_heading(doc, "6. Жесты Appium (дополнительно)")
    add_para(
        doc,
        "Это не SauceDemo, а ApiDemos — просто чтобы закрыть часть про свайпы "
        "и scroll. Результат прошлого прогона оставляю как есть.",
    )
    add_table(doc, ["ID", "Сценарий", "Результат"], GESTURE_RESULTS)
    add_code(doc, "============================= 5 passed in 54.93s ==============================")

    add_heading(doc, "7. Ручное тестирование")
    add_para(
        doc,
        "Человеческий checklist я не проходил. Статус: PENDING HUMAN VERIFICATION "
        "(файл mobile/docs/kt09-manual-checklist.md). Автоматические скриншоты "
        "за ручные не выдаю — это разные вещи.",
        bold=True,
    )

    add_heading(doc, "8. Что заметил по ходу")
    add_para(
        doc,
        "• По функционалу SauceDemo на Windows/Android automation явных багов "
        "сайта не поймал.",
    )
    add_para(
        doc,
        "• На эмуляторе Android Chrome иногда всплывает Google Password Manager, "
        "а на части кадров — полосы от soft GPU. Это окружение, не дефект Swag Labs.",
    )
    add_para(doc, "• iOS не трогал — NOT RUN.")

    add_heading(doc, "9. Заключение")
    add_para(
        doc,
        "КТ 09: SauceDemo прогнан автоматизацией на Windows Firefox и Android Chrome "
        "(emulator-5554, Chrome 134, Appium 3.8), по восемь скриншотов на сторону. "
        "Сохранены прошлые цифры: 11 passed in 36.91s (Android DOM), "
        "5 passed in 54.93s (жесты), 2 passed (Windows history). "
        "Ручная проверка: PENDING HUMAN VERIFICATION.",
    )

    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    required = [
        "КТ №9",
        "saucedemo.com",
        "AUTOMATED",
        "PENDING HUMAN VERIFICATION",
        "11 passed in 36.91s",
        "5 passed in 54.93s",
        "Заключение",
        "Windows",
        "Android",
    ]
    missing = [r for r in required if r not in text]
    if missing:
        raise AssertionError(missing)
    pictures = list(doc.inline_shapes)
    if len(pictures) < 14:
        raise AssertionError(f"pictures={len(pictures)}")
    images = [r for r in doc.part.rels.values() if r.reltype == RT.IMAGE]
    if not images:
        raise AssertionError("no images")
    print(f"VERIFY_OK size={path.stat().st_size} pictures={len(pictures)} unique={len(images)}")


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    verify(out)
