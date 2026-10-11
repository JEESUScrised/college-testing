"""Generate reports/KT01.docx from verified KT01 artifacts."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SCREENSHOT = ROOT / "screenshots" / "kt01" / "ya_ru_opened.png"
OUTPUT = Path(__file__).resolve().parent / "KT01.docx"

TEST_OUTPUT = """\
============================= test session starts =============================
platform win32 -- Python 3.13.9, pytest-8.4.2, pluggy-1.6.0 -- ...\\.venv\\Scripts\\python.exe
cachedir: .pytest_cache
rootdir: ...\\college-testing
configfile: pytest.ini
collecting ... collected 1 item

selenium/tests/test_kt01_ya_ru.py::test_open_ya_ru PASSED                [100%]

============================== 1 passed in 4.99s ==============================
"""

FIXTURE_CODE = """\
def pytest_addoption(parser):
    parser.addoption("--browser", default="chrome", choices=("chrome", "firefox"))

@pytest.fixture
def driver(browser_name):
    options = ChromeOptions()
    options.add_argument("--window-size=1280,900")
    drv = webdriver.Chrome(options=options)
    drv.set_page_load_timeout(30)
    drv.implicitly_wait(5)
    drv.set_script_timeout(30)
    try:
        yield drv
    finally:
        drv.quit()
"""

TEST_CODE = """\
@pytest.mark.kt01
def test_open_ya_ru(driver, save_screenshot):
    target = "https://ya.ru"
    driver.get(target)
    current = driver.current_url
    host = urlparse(current).netloc.lower()
    title = (driver.title or "").strip()
    assert host.endswith("ya.ru") or "yandex" in host
    assert title
    assert driver.execute_script("return document.readyState") == "complete"
    screenshot_path = save_screenshot("kt01", "ya_ru_opened")
    assert screenshot_path.is_file() and screenshot_path.stat().st_size > 0
"""


def build() -> Path:
    if not SCREENSHOT.is_file():
        raise FileNotFoundError(f"Screenshot not found: {SCREENSHOT}")

    doc = Document()
    setup_a4(doc)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(doc, "КТ №1. Запуск и работа с WebDriver в Selenium", bold=True, center=True)
    add_para(doc, "Отчёт по контрольной точке", center=True)
    add_para(doc, "Дата выполнения: 2026-10-09", center=True)

    add_heading(doc, "Зачем я это делал")
    add_para(
        doc,
        "Первая КТ — просто поднять Selenium WebDriver 4 и убедиться, что браузер "
        "реально открывается из Python. Задача небольшая: Chrome, переход на "
        "https://ya.ru, пара assertions и скриншот. Без Page Object и сложных сценариев — "
        "сначала проверить, что связка pytest + Selenium вообще живая.",
    )

    add_heading(doc, "Что требовалось по заданию")
    add_para(doc, "КТ 01, 5 баллов. По сути:")
    add_para(doc, "• автотест на Python (pytest + Selenium 4);")
    add_para(doc, "• запуск Google Chrome через WebDriver;")
    add_para(doc, "• открыть https://ya.ru;")
    add_para(doc, "• проверить переход нормальными assertions;")
    add_para(doc, "• сохранить скриншот и оформить отчёт.")

    add_heading(doc, "На чём гонял")
    add_para(doc, "• ОС: Windows 11 (10.0.26100)")
    add_para(doc, "• Python 3.13.9")
    add_para(doc, "• pytest 8.4.2")
    add_para(doc, "• Selenium 4.50.0")
    add_para(doc, "• Google Chrome 154.0.8037.98")
    add_para(doc, "• драйвер подтянул Selenium Manager (идёт с Selenium 4)")
    add_para(doc, "• URL: https://ya.ru")
    add_para(doc, "• код лежит в selenium/")

    add_heading(doc, "Как ставил и запускал")
    add_para(doc, "Сначала venv и зависимости:")
    add_code(
        doc,
        """py -m venv .venv
.\\.venv\\Scripts\\Activate.ps1
python -m pip install --upgrade pip
pip install -r selenium\\requirements.txt""",
    )
    add_para(doc, "Сам тест КТ 01:")
    add_code(doc, "pytest selenium\\tests\\test_kt01_ya_ru.py -v --browser=chrome")
    add_para(
        doc,
        "Через --browser=chrome фикстура выбирает браузер. Скриншот пишется в "
        "screenshots/kt01/ya_ru_opened.png.",
    )

    add_heading(doc, "Фикстура и что именно проверяю")
    add_para(
        doc,
        "Фикстура driver — в selenium/conftest.py. Читает --browser, по умолчанию "
        "стартует Chrome, выставляет page load 30 с, implicit wait 5 с, script timeout "
        "30 с, в finally зовёт quit(). save_screenshot кладёт PNG в screenshots/<кт>/.",
    )
    add_para(doc, "Кусок фикстуры:")
    add_code(doc, FIXTURE_CODE)
    add_para(
        doc,
        "test_open_ya_ru идёт на https://ya.ru и смотрит четыре вещи: host — ya.ru или "
        "yandex (редирект бывает), title не пустой, document.readyState == complete, "
        "скриншот существует и размер > 0.",
    )
    add_para(doc, "Кусок теста:")
    add_code(doc, TEST_CODE)

    add_heading(doc, "Что получилось на прогоне")
    add_para(
        doc,
        "После того как добавил pytest.ini (маркер kt01), прогнал ещё раз 2026-10-09. "
        "Предупреждений не было. Результат — PASSED.",
    )
    add_para(doc, "Вывод терминала:")
    add_code(doc, TEST_OUTPUT)
    add_para(doc, "Итог: 1 passed in 4.99s. Exit code: 0.")

    add_heading(doc, "Скриншот")
    add_para(
        doc,
        "Скриншот снял сам WebDriver после открытия ya.ru — файл "
        "screenshots/kt01/ya_ru_opened.png.",
    )
    add_shot(doc, SCREENSHOT, "Рисунок 1 — Главная страница ya.ru после запуска WebDriver", width_cm=15.0)

    add_heading(doc, "Заключение")
    add_para(
        doc,
        "Для первой КТ этого хватило: venv, Selenium 4, фикстура с --browser=chrome, "
        "один тест на https://ya.ru. Прогон дал 1 passed in 4.99s, скриншот "
        "ya_ru_opened.png на месте. Дальше уже можно наращивать сценарии — база для "
        "WebDriver у меня есть.",
    )
    add_para(
        doc,
        "Минус: ya.ru может менять вёрстку и показывать баннеры. Мой тест смотрит только "
        "на факт загрузки, а не на поиск или весь UI.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    required = [
        "КТ №1. Запуск и работа с WebDriver в Selenium",
        "1 passed in 4.99s",
        "ya.ru",
        "ya_ru_opened.png",
        "Заключение",
    ]
    missing = [item for item in required if item not in text]
    if missing:
        raise AssertionError(f"DOCX missing sections: {missing}")
    if len(list(doc.inline_shapes)) < 1:
        raise AssertionError("Expected at least 1 picture")
    if not any(rel.reltype == RT.IMAGE for rel in doc.part.rels.values()):
        raise AssertionError("No embedded image parts found")
    print(f"VERIFY_OK size={path.stat().st_size} pictures={len(list(doc.inline_shapes))}")


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    verify(out)
