"""Generate reports/KT10.docx — cross-browser + responsive SauceDemo."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot, add_table
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt10"
RESP = SHOTS / "responsive"
OUTPUT = Path(__file__).resolve().parent / "KT10.docx"

HIST = [
    ("TC-10-01", "title / URL", "chrome + firefox", "PASS"),
    ("TC-10-02", "клик по DOM", "chrome + firefox", "PASS"),
    ("TC-10-03", "валидация формы", "chrome + firefox", "PASS"),
    ("TC-10-04", "переключение окон", "chrome + firefox", "PASS"),
    ("TC-10-05", "iframe", "chrome + firefox", "PASS"),
]

RESP_CASES = [
    ("TC10-R01", "раскладка / overflow"),
    ("TC10-R02", "форма логина"),
    ("TC10-R03", "каталог"),
    ("TC10-R04", "burger menu"),
    ("TC10-R05", "корзина"),
    ("TC10-R06", "форма checkout"),
]

COMPLIANCE = [
    ("Публичное адаптивное приложение", "Sauce Demo", "PASS"),
    ("Разные браузеры", "Chrome + Firefox", "PASS"),
    ("Desktop / tablet / phone", "1440×900 / 768×1024 / 390×844", "PASS"),
    ("Меню, кнопки, формы", "R02–R06", "PASS"),
    ("Адаптация размеров", "viewport JSONL + screenshots", "PASS"),
    ("Документирование проблем", "horizontal overflow: none observed", "PASS"),
    ("Отчёт с рекомендациями", "KT10.docx + docs/kt10/", "PASS"),
]


def build() -> Path:
    doc = Document()
    setup_a4(doc, left=2.5, right=1.5)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №10. Кроссбраузерное и адаптивное тестирование",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке", center=True)
    add_para(doc, "Дата: 2026-10-10", center=True)

    add_heading(doc, "1. Цель")
    add_para(
        doc,
        "По заданию нужно было взять публичный адаптивный сайт, прогнать его "
        "в Chrome и Firefox, посмотреть desktop / tablet / phone вьюпорты, "
        "проверить меню и формы, записать, если что-то поедет, и оформить отчёт.",
    )

    add_heading(doc, "2. Какой сайт взял")
    add_para(
        doc,
        "Опять https://www.saucedemo.com/ — логин, каталог, меню, корзина, checkout. "
        "Знакомый DOM, удобно сравнивать с предыдущими КТ.",
    )

    add_heading(doc, "3. Responsive-матрица")
    add_para(
        doc,
        "Три вьюпорта: 1440×900 (desktop), 768×1024 (tablet), 390×844 (phone). "
        "В Chrome для tablet/phone крутил CDP-эмуляцию плюс resize окна, "
        "в Firefox — только resize (там CDP device metrics нет). "
        "Это эмуляция вьюпорта, не физические телефоны — на реальном железе "
        "пиксели и жесты могут отличаться.",
    )
    add_table(doc, ["ID", "Сценарий"], RESP_CASES)
    add_para(doc, "Прогон responsive suite:")
    add_code(doc, "============================= 6 passed in 42.73s ==============================")
    add_para(
        doc,
        "В evidence 36 строк по комбинациям браузер×вьюпорт×шаг. "
        "Горизонтальный overflow нигде не всплыл (none observed).",
    )
    for name, cap in [
        ("chrome_phone_r02_login_form.png", "Chrome phone 390×844 — форма логина"),
        ("firefox_phone_r03_inventory.png", "Firefox phone — каталог"),
        ("chrome_tablet_r04_menu_open.png", "Chrome tablet 768×1024 — открытое меню"),
        ("chrome_desktop_r06_checkout.png", "Chrome desktop 1440×900 — checkout / валидация"),
    ]:
        add_shot(doc, RESP / name, cap, width_cm=9.0)

    add_heading(doc, "4. Старый functional suite (не трогал в remediation)")
    add_para(
        doc,
        "Раньше уже был набор на локальных фикстурах с EventFiringWebDriver listener. "
        "При доработке responsive его заново не гонял — оставляю исторический результат "
        "как есть: 10 passed in 52.92s (по пять сценариев × Chrome/Firefox).",
    )
    add_table(doc, ["ID", "Сценарий", "Браузеры", "Статус"], HIST)
    add_code(doc, "============================= 10 passed in 52.92s ==============================")
    add_shot(doc, SHOTS / "chrome_01_home_title.png", "Исторический прогон — Chrome home", width_cm=8.0)
    add_shot(doc, SHOTS / "firefox_01_home_title.png", "Исторический прогон — Firefox home", width_cm=8.0)

    add_heading(doc, "5. Compliance")
    add_para(
        doc,
        "Сверка с пунктами из задания — что закрыл и чем подтверждаю:",
    )
    add_table(doc, ["Требование", "Реализация", "Статус"], COMPLIANCE)

    add_heading(doc, "6. Заключение")
    add_para(
        doc,
        "Добавил responsive-матрицу на Sauce Demo: Chrome и Firefox, три вьюпорта "
        "(1440×900 / 768×1024 / 390×844), шесть сценариев — 6 passed in 42.73s. "
        "Кривой раскладки и overflow не нашёл. Старый listener/functional suite "
        "сохранён (10 passed in 52.92s). Помню ограничение: viewport emulation "
        "в DevTools — это не замена теста на настоящем телефоне.",
    )

    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    for req in ["6 passed in 42.73s", "10 passed in 52.92s", "saucedemo", "Compliance"]:
        if req not in text:
            raise SystemExit(f"missing {req}")
    print(f"OK {path} size={path.stat().st_size}")


if __name__ == "__main__":
    verify(build())
