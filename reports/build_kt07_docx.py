"""Generate reports/KT07.docx from verified Selenium Grid evidence (Russian)."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot, add_table, set_run_font
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt07"
MULTI = SHOTS / "multi"
OUTPUT = Path(__file__).resolve().parent / "KT07.docx"

SCREENSHOTS = [
    ("01_remote_session_blank.png", "Рисунок 1 — Remote-сессия через Grid"),
    ("02_navigate_title.png", "Рисунок 2 — Навигация и title"),
    ("03_dom_interaction.png", "Рисунок 3 — Взаимодействие с DOM"),
    ("04_windows_main.png", "Рисунок 4 — Главное окно"),
    ("05_windows_secondary.png", "Рисунок 5 — Вторичное окно"),
    ("06_windows_back_main.png", "Рисунок 6 — Возврат в главное окно"),
    ("07_iframe_host.png", "Рисунок 7 — Host с iframe"),
    ("08_iframe_inside.png", "Рисунок 8 — Внутри iframe"),
    ("09_iframe_default_content.png", "Рисунок 9 — default_content"),
]

MULTI_SHOTS = [
    (MULTI / "chrome_01_session.png", "Рисунок 10 — Multi-node Chrome session"),
    (MULTI / "firefox_01_session.png", "Рисунок 11 — Multi-node Firefox session"),
    (MULTI / "chrome_02_landing.png", "Рисунок 12 — Chrome fixture via Grid"),
    (MULTI / "firefox_02_landing.png", "Рисунок 13 — Firefox fixture via Grid"),
    (MULTI / "parallel_chrome_overlap.png", "Рисунок 14 — Parallel Chrome (overlap proof)"),
    (MULTI / "parallel_firefox_overlap.png", "Рисунок 15 — Parallel Firefox (overlap proof)"),
]

RESULTS = [
    ("TC-07-01", "test_remote_session_created", "PASS"),
    ("TC-07-02", "test_navigate_and_title", "PASS"),
    ("TC-07-03", "test_dom_element_interaction", "PASS"),
    ("TC-07-04", "test_window_open_and_switch", "PASS"),
    ("TC-07-05", "test_iframe_interaction_and_context_switch", "PASS"),
]


def build() -> Path:
    doc = Document()
    setup_a4(doc)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №7. Удалённый запуск автотестов через Selenium Grid 4",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке", center=True)
    add_para(doc, "Дата: 2026-10-09", center=True)

    add_heading(doc, "1. Цель и требования")
    add_para(
        doc,
        "Нужно было поднять Selenium Grid 4 у себя на машине, гонять браузерные "
        "тесты через RemoteWebDriver и показать, что работают и Standalone, "
        "и схема Hub + Node. Без реальных сессий, логов и скриншотов отчёт "
        "бы не закрыл.",
    )

    add_heading(doc, "2. Архитектура Selenium Grid")
    add_para(
        doc,
        "В Standalone Router, Distributor, Session Map, Queue и Node живут "
        "в одном процессе — удобно для старта. В режиме Hub + Node точка входа "
        "отдельная (Hub), а браузеры крутит Node. Клиент в обоих случаях "
        "ходит на http://127.0.0.1:4444.",
    )
    add_para(
        doc,
        "Важный момент: RemoteWebDriver на localhost — это всё равно удалённый "
        "протокол Grid, даже если Hub и браузер на одном ПК. Кластер из "
        "нескольких машин я не поднимал и на мультимашинность не претендую.",
        bold=True,
    )

    add_heading(doc, "3. Среда")
    add_para(doc, "Windows 10 x64; Java 21 (Android Studio JBR); Python 3.13.2; Selenium 4.50.0.")
    add_para(doc, "Chrome 155.0.8059.39; selenium-server-4.50.0.jar (SeleniumHQ GitHub Releases).")
    add_para(doc, "Grid bind: 127.0.0.1. Node port: 5556 (5555 занят Android emulator/qemu).")
    add_para(
        doc,
        "Event Bus у Hub и Node явно повесил на tcp://127.0.0.1:4442 и :4443. "
        "Без этого Grid начинал светить VPN-IP (26.x), и Node просто "
        "не регистрировался — долго ловил этот момент.",
    )

    add_heading(doc, "4. Команды запуска")
    add_code(
        doc,
        r""".\selenium\grid\download_server.ps1
.\selenium\grid\start_standalone.ps1
.\selenium\grid\health_check.ps1
.\scripts\run_kt07.ps1
.\selenium\grid\stop_grid.ps1

# Hub + Node (не одновременно со Standalone на :4444)
.\selenium\grid\start_hub_node.ps1
.\selenium\grid\health_check.ps1
.\selenium\grid\stop_grid.ps1""",
    )

    add_heading(doc, "5. Remote WebDriver")
    add_para(
        doc,
        "Фикстура grid_driver создаёт webdriver.Remote(command_executor=grid_url, "
        "options=ChromeOptions()) и не трогает локальный driver из KT01–KT05. "
        "HTML-фикстуры отдаёт обычный HTTP-сервер, не file:// — иначе Grid "
        "иногда капризничает.",
    )
    add_code(
        doc,
        'options = ChromeOptions()\n'
        'options.set_capability("se:name", "KT07:...")\n'
        'driver = webdriver.Remote("http://127.0.0.1:4444", options=options)\n'
        "# session_id / capabilities сохраняются в selenium/artifacts/kt07/",
    )

    add_heading(doc, "6. Результаты Standalone suite")
    add_para(
        doc,
        "Диагностику начал с одного кейса — PASS. Потом полный suite: "
        "5 passed in 4.88s, exit 0.",
    )
    add_table(doc, ["ID", "Автотест", "Результат"], RESULTS)
    add_code(doc, "============================== 5 passed in 4.88s ==============================")

    add_heading(doc, "7. Hub + Node")
    add_para(
        doc,
        "Standalone остановил и поднял Hub на :4444 плюс Node на :5556. "
        "Статус: ready=true, nodes=1, uri Node = http://127.0.0.1:5556. "
        "Что сессии реально идут через Node, проверил двумя тестами "
        "(session + navigate) — оба PASS.",
    )

    add_heading(doc, "8. Remediation: Hub + 2 Nodes + parallel (2026-10-10)")
    add_para(
        doc,
        "Дальше добавил режим Hub + Node A (:5556) + Node B (:5557). "
        "В статусе Grid было ready=true, nodes=2. Multi-browser suite "
        "Chrome+Firefox: 4 passed in 7.72s. Параллельный proof: одновременно "
        "держались Chrome и Firefox Remote-сессии, интервалы hold пересеклись "
        "(parallel_overlap=true). Это как раз и показывает, что Grid раздаёт "
        "сессии параллельно, а не «по очереди втихую».",
    )
    add_code(
        doc,
        r""".\selenium\grid\start_hub_multi_node.ps1
.\.venv\Scripts\python.exe -m pytest selenium\tests\test_kt07_multi_node.py -v --grid-url=http://127.0.0.1:4444
.\.venv\Scripts\python.exe scripts\run_kt07_parallel_proof.py
.\selenium\grid\stop_grid.ps1""",
    )
    add_code(
        doc,
        "============================== 4 passed in 7.72s ==============================\n"
        "parallel_overlap=true (Chrome hold ∩ Firefox hold)",
    )

    add_heading(doc, "9. Скриншоты — исторический Standalone")
    for name, caption in SCREENSHOTS:
        add_shot(doc, SHOTS / name, caption, width_cm=12.0)

    add_heading(doc, "10. Скриншоты — multi-node / parallel")
    for path, caption in MULTI_SHOTS:
        add_shot(doc, path, caption, width_cm=12.0)

    add_heading(doc, "11. Заключение")
    add_para(
        doc,
        "По КТ 07 Standalone остаётся 5 passed in 4.88s. На remediation "
        "добавил Hub + два Node, прогнал Chrome+Firefox (4 passed in 7.72s) "
        "и зафиксировал parallel_overlap. Весь Grid сидел на 127.0.0.1 — "
        "это локальный стенд, не ферма из нескольких ПК. "
        "Главный практический вывод: без явного Event Bus на localhost Node "
        "легко «теряется» из-за VPN-интерфейса, и это лучше проверить "
        "health_check до прогона suite.",
    )

    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    required = [
        "КТ №7. Удалённый запуск автотестов через Selenium Grid 4",
        "5 passed in 4.88s",
        "4 passed in 7.72s",
        "parallel_overlap",
        "127.0.0.1",
        "Заключение",
    ]
    missing = [r for r in required if r not in text]
    if missing:
        raise AssertionError(missing)
    pictures = list(doc.inline_shapes)
    expected = len(SCREENSHOTS) + len(MULTI_SHOTS)
    if len(pictures) < expected:
        raise AssertionError(f"pictures={len(pictures)} expected>={expected}")
    print(f"VERIFY_OK size={path.stat().st_size} pictures={len(pictures)}")


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    verify(out)
