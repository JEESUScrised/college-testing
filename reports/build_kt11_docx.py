"""Generate reports/KT11.docx — Robot Framework keyword-driven suite."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt11"
OUTPUT = Path(__file__).resolve().parent / "KT11.docx"

RESULTS = [
    ("TC-11-01", "страница открывается / title", "PASS", "2.74"),
    ("TC-11-02", "навигация по ссылке", "PASS", "2.80"),
    ("TC-11-03", "валидная отправка формы", "PASS", "2.81"),
    ("TC-11-04", "невалидная форма отклоняется", "PASS", "2.81"),
    ("TC-11-05", "checkbox вкл/выкл", "PASS", "2.78"),
    ("TC-11-06", "выбор radio", "PASS", "2.79"),
    ("TC-11-07", "dropdown", "PASS", "2.80"),
    ("TC-11-08", "iframe: действие и возврат", "PASS", "2.82"),
    ("TC-11-09", "новая вкладка и возврат", "PASS", "2.88"),
    ("TC-11-10", "динамическое обновление UI", "PASS", "2.77"),
]

IMAGES = [
    ("01_home_title.png", "Рисунок 1 — Home / title"),
    ("02_navigation_register.png", "Рисунок 2 — Переход на регистрацию"),
    ("03_form_valid.png", "Рисунок 3 — Валидная форма принята"),
    ("04_form_invalid.png", "Рисунок 4 — Ошибка валидации"),
    ("05_checkbox.png", "Рисунок 5 — Checkbox"),
    ("06_radio.png", "Рисунок 6 — Radio"),
    ("07_dropdown.png", "Рисунок 7 — Dropdown"),
    ("08_iframe.png", "Рисунок 8 — IFrame после взаимодействия"),
    ("09_secondary_tab.png", "Рисунок 9 — Вторая вкладка"),
    ("10_dynamic_update.png", "Рисунок 10 — Динамический статус"),
    ("11_robot_report.png", "Рисунок 11 — Robot report.html (10 passed)"),
    ("12_robot_log.png", "Рисунок 12 — Robot log.html"),
]

ROBOT_SAMPLE = r"""*** Test Cases ***
TC-11-03 Valid Form Submission Succeeds
    [Tags]    KT11    form    positive
    Open Registration Page
    Submit Registration Form    Анна Тестова    anna.kt11@example.com
    Verify Form Accepted    Анна Тестова    anna.kt11@example.com
    Capture Scenario Screenshot    03_form_valid"""


def build() -> Path:
    doc = Document()
    setup_a4(doc, left=2.5, right=1.5)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №11. Автоматизированное тестирование с Robot Framework",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке", center=True)
    add_para(doc, "Дата: 2026-10-10", center=True)

    add_heading(doc, "1. Что требовалось")
    add_para(
        doc,
        "Сделать ровно десять разных тест-кейсов на «родном» синтаксисе Robot Framework, "
        "подключить SeleniumLibrary, разложить шаги по keywords, поднять локальное "
        "приложение и сдать стандартные отчёты Robot (report/log/output).",
    )
    add_para(
        doc,
        "Слайды cloud.ithub.ru по КТ11 без логина не открываются, поэтому "
        "на «один в один со скрытой презентацией» не претендую — опирался "
        "на формулировку задания и здравый смысл.",
        bold=True,
    )

    add_heading(doc, "2. Коротко про Robot Framework")
    add_para(
        doc,
        "Robot — это keyword-driven штука: сценарии пишешь таблицами в .robot, "
        "а переиспользуемые шаги выносишь в resource-файлы. Секции обычные — "
        "Settings, Variables, Test Cases, Keywords. После прогона получаешь "
        "output.xml, подробный log.html и сводный report.html.",
    )

    add_heading(doc, "3. SeleniumLibrary и как разложил keywords")
    add_para(
        doc,
        "SeleniumLibrary даёт низкоуровневые слова вроде Open Browser, Click Element, "
        "Select Frame, Switch Window. Я поверх них собрал свои: Open Test Application, "
        "Submit Registration Form, Verify Form Validation, Open Secondary Tab, "
        "Return To Main Page, Interact Inside Demo Iframe — чтобы в тест-кейсах "
        "читалось «что делаю», а не xpath.",
    )
    add_para(
        doc,
        "Python-хелпер FixtureServer.py только крутит локальный HTTP на свободном порту. "
        "Сами проверки живут в .robot, логику в питон не утаскивал.",
    )

    add_heading(doc, "4. Окружение")
    add_para(
        doc,
        "Windows 10, Python 3.13.2, Robot Framework 7.5, SeleniumLibrary 6.9.0, "
        "Selenium 4.50.0, Chrome. Зависимости в robot/requirements.txt — selenium "
        "специально зафиксировал на 4.50.0, чтобы не разъехались КТ01–КТ10.",
    )

    add_heading(doc, "5. На чём гонял")
    add_para(
        doc,
        "Локальные HTML из robot/fixtures/: навигация, форма, checkbox/radio/dropdown, "
        "iframe, вторая вкладка, динамический статус. Без чужих сайтов и VPN — "
        "так стабильнее и не зависит от капчи.",
    )

    add_heading(doc, "6. Десять тест-кейсов")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "Сценарий", "Статус", "с"
    for row in RESULTS:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    add_para(doc, "")

    add_heading(doc, "7. Кусок .robot")
    add_para(doc, "Пример позитивного кейса с формой — видно, как вызываются keywords:")
    add_code(doc, ROBOT_SAMPLE)

    add_heading(doc, "8. Resources и как запускал")
    add_para(
        doc,
        "common.resource — bootstrap/teardown сьюта, Open/Close Test Application, "
        "скриншоты. pages.resource — страничные keywords (формы, iframe, вкладки, dynamic).",
    )
    add_code(
        doc,
        r""".\scripts\run_kt11.ps1
.\.venv\Scripts\python.exe -m robot --outputdir robot\artifacts\kt11 robot\tests\kt11.robot""",
    )

    add_heading(doc, "9. Фактический прогон")
    add_para(doc, "Dry-run нашёл 10 тестов. Отдельно прогнал TC-11-01 — PASS.")
    add_para(
        doc,
        "Полный suite: 10 tests, 10 passed, 0 failed; по времени вышло ≈ 28.39 с.",
    )
    add_code(doc, "10 tests, 10 passed, 0 failed\nSUITE PASS elapsed= 28.389725")
    add_para(
        doc,
        "Артефакты лежат в robot/artifacts/kt11/: output.xml, log.html, report.html, xunit.xml.",
    )

    add_heading(doc, "10. Скриншоты сценариев и отчётов Robot")
    add_para(
        doc,
        "Десять кадров по шагам плюс report.html и log.html — всего 12 картинок.",
    )
    for filename, caption in IMAGES:
        width = 14.0 if filename.startswith("11_") or filename.startswith("12_") else 7.5
        add_shot(doc, SHOTS / filename, caption, width_cm=width)

    add_heading(doc, "11. Заключение")
    add_para(
        doc,
        "КТ 11 закрыл: десять native Robot-тестов, SeleniumLibrary, keyword-driven "
        "resources, локальные фикстуры. Suite прошёл 10/10 за ~28.39 с, отчёты Robot "
        "и этот DOCX на месте. КТ12 отдельно.",
    )

    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    required = [
        "КТ №11. Автоматизированное тестирование с Robot Framework",
        "10 tests, 10 passed, 0 failed",
        "SeleniumLibrary",
        "Заключение",
    ]
    missing = [r for r in required if r not in text]
    if missing:
        raise SystemExit(f"DOCX missing: {missing}")
    n_pics = sum(1 for r in doc.part.rels.values() if "image" in r.reltype)
    if n_pics < 10:
        raise SystemExit(f"Expected >=10 images, got {n_pics}")
    print(f"OK {path} images={n_pics} chars={len(text)}")


if __name__ == "__main__":
    out = build()
    verify(out)
