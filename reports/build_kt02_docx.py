"""Generate reports/KT02.docx from verified KT02 artifacts."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt02"
OUTPUT = Path(__file__).resolve().parent / "KT02.docx"

SCREENSHOTS = [
    ("01_main_window_before_open.png", "Рисунок 1 — Главное окно до открытия новой вкладки"),
    ("03_switched_to_new_window.png", "Рисунок 2 — Переключение на новое окно по window handle"),
    ("04_new_window_content_verified.png", "Рисунок 3 — Проверка содержимого вторичного окна"),
    ("06_returned_to_main_window.png", "Рисунок 4 — Возврат в исходное окно после закрытия вторичного"),
    ("07_iframe_host_before_switch.png", "Рисунок 5 — Основной документ с iframe"),
    ("08_iframe_interaction_result.png", "Рисунок 6 — Взаимодействие внутри iframe и проверка результата"),
    ("10_back_to_default_content.png", "Рисунок 7 — Возврат в основной документ через default_content()"),
]

KT02_OUTPUT = """\
============================= test session starts =============================
platform win32 -- Python 3.13.9, pytest-8.4.2, pluggy-1.6.0
configfile: pytest.ini
collecting ... collected 6 items

selenium/tests/test_kt02_windows.py::test_open_new_browser_window PASSED
selenium/tests/test_kt02_windows.py::test_switch_to_new_window_by_handle PASSED
selenium/tests/test_kt02_windows.py::test_verify_new_window_content PASSED
selenium/tests/test_kt02_windows.py::test_close_secondary_and_return_to_original PASSED
selenium/tests/test_kt02_iframe.py::test_switch_into_iframe_interact_and_verify PASSED
selenium/tests/test_kt02_iframe.py::test_return_to_default_content PASSED

============================= 6 passed in 21.49s ==============================
"""

KT01_REGRESSION = """\
selenium/tests/test_kt01_ya_ru.py::test_open_ya_ru PASSED
============================== 1 passed in 4.53s ==============================
"""

WINDOWS_CODE = """\
new_handle = _open_secondary_window(driver, wait, windows_main)
driver.switch_to.window(new_handle)
heading = wait.until(EC.visibility_of_element_located((By.ID, "new-heading")))
assert heading.text.strip() == "Содержимое нового окна"
driver.close()
wait.until(EC.number_of_windows_to_be(1))
driver.switch_to.window(windows_main)
"""

IFRAME_CODE = """\
driver.switch_to.frame(frame)
field.send_keys("КТ02 iframe OK")
submit.click()
wait.until(EC.text_to_be_present_in_element((By.ID, "frame-result"), "КТ02 iframe OK"))
driver.switch_to.default_content()
assert host_heading.text.strip() == "Основной документ с iframe"
"""


def set_run_font(run, size_pt: float = 12, bold: bool = False, name: str = "Times New Roman") -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size_pt=14 if level == 1 else 13, bold=True)


def add_para(doc: Document, text: str, *, bold: bool = False, center: bool = False) -> None:
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15


def add_code(doc: Document, code: str) -> None:
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        set_run_font(run, size_pt=9, name="Consolas")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0


def add_shot(doc: Document, filename: str, caption: str) -> None:
    path = SHOTS / filename
    if not path.is_file():
        raise FileNotFoundError(path)
    doc.add_picture(str(path), width=Cm(14))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    set_run_font(run, size_pt=11)


def build() -> Path:
    missing = [name for name, _ in SCREENSHOTS if not (SHOTS / name).is_file()]
    if missing:
        raise FileNotFoundError(f"Missing screenshots: {missing}")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    add_para(doc, "Дисциплина: Модульное тестирование веб-приложений", center=True)
    add_para(doc, "Учебное заведение: ______________________________", center=True)
    add_para(doc, "Студент: _______________________________________", center=True)
    add_para(doc, "Преподаватель: _________________________________", center=True)
    add_para(doc, "Группа: ________________________________________", center=True)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №2. Тестирование функционала веб-приложения с применением техник тестирования",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке", center=True)
    add_para(doc, "Дата выполнения: 2026-10-09", center=True)

    add_heading(doc, "1. Цель работы")
    add_para(
        doc,
        "Цель контрольной точки №2 — освоить техники тестирования веб-приложения "
        "с помощью Selenium WebDriver: работу с несколькими окнами/вкладками "
        "(window handles) и взаимодействие с содержимым iframe, включая возврат "
        "в основной документ методом default_content().",
    )

    add_heading(doc, "2. Требования задания")
    add_para(doc, "Согласно постановке КТ 02 (10 баллов) и уточнённым требованиям:")
    add_para(doc, "• открыть новое окно/вкладку браузера;")
    add_para(doc, "• переключиться на новое окно через window handles;")
    add_para(doc, "• проверить содержимое нового окна;")
    add_para(doc, "• закрыть вторичное окно и вернуться в исходное;")
    add_para(doc, "• переключиться в iframe, выполнить действие и проверить результат;")
    add_para(doc, "• вернуться в основной документ через default_content();")
    add_para(doc, "• использовать осмысленные assertions и явные ожидания WebDriverWait;")
    add_para(doc, "• сохранить инфраструктуру КТ 01 без поломки регрессии.")

    add_heading(doc, "3. Выбор тестового приложения")
    add_para(
        doc,
        "Публичный демонстрационный сайт https://the-internet.herokuapp.com был "
        "проверен перед реализацией: страницы /windows и /iframe отвечали HTTP 200. "
        "Для стабильных локаторов и воспроизводимости без зависимости от внешнего "
        "TinyMCE-редактора подготовлены локальные HTML-фикстуры в "
        "selenium/fixtures/kt02/ (аналог учебных страниц the-internet: отдельное "
        "окно и вложенный iframe с формой).",
    )

    add_heading(doc, "4. Программное обеспечение и среда")
    add_para(doc, "• ОС: Windows 11 (10.0.26100)")
    add_para(doc, "• Python 3.13.9, pytest 8.4.2, Selenium 4.50.0")
    add_para(doc, "• Google Chrome 154.0.8037.98")
    add_para(doc, "• Переиспользуемая фикстура driver / --browser=chrome из КТ 01")
    add_para(doc, "• Явные ожидания: WebDriverWait (10 с), EC.*")

    add_heading(doc, "5. Установка и запуск")
    add_code(
        doc,
        """.\\.venv\\Scripts\\Activate.ps1
pip install -r selenium\\requirements.txt
pytest selenium\\tests\\test_kt02_windows.py selenium\\tests\\test_kt02_iframe.py -v --browser=chrome
pytest selenium\\tests\\test_kt01_ya_ru.py -v --browser=chrome""",
    )

    add_heading(doc, "6. Тест-кейсы")
    add_para(doc, "TC-W1. test_open_new_browser_window — открытие второго окна по ссылке target=_blank.")
    add_para(doc, "TC-W2. test_switch_to_new_window_by_handle — переключение driver.switch_to.window(handle).")
    add_para(doc, "TC-W3. test_verify_new_window_content — проверка заголовка и текста вторичного окна.")
    add_para(doc, "TC-W4. test_close_secondary_and_return_to_original — close() и возврат в исходный handle.")
    add_para(doc, "TC-F1. test_switch_into_iframe_interact_and_verify — frame(), ввод текста, проверка результата.")
    add_para(doc, "TC-F2. test_return_to_default_content — default_content() и проверка элементов host-страницы.")

    add_heading(doc, "7. Пояснения к реализации")
    add_para(
        doc,
        "Инфраструктура КТ 01 сохранена. В conftest.py добавлены вспомогательные "
        "фикстуры wait и fixture_url без изменения поведения KT01. Открытие второго "
        "окна ожидается через EC.number_of_windows_to_be; переключение выполняется "
        "по разнице множеств window_handles. Для iframe используется "
        "EC.frame_to_be_available_and_switch_to_it / switch_to.frame, взаимодействие "
        "с input/button и проверка текста результата; возврат — switch_to.default_content().",
    )
    add_para(doc, "Фрагмент сценария окон:")
    add_code(doc, WINDOWS_CODE)
    add_para(doc, "Фрагмент сценария iframe:")
    add_code(doc, IFRAME_CODE)

    add_heading(doc, "8. Фактические результаты")
    add_para(doc, "Прогон КТ 02 (2026-10-09):")
    add_code(doc, KT02_OUTPUT)
    add_para(doc, "Регрессия КТ 01:")
    add_code(doc, KT01_REGRESSION)
    add_para(doc, "Итог: KT02 — 6 passed за 21.49 с; KT01 — 1 passed за 4.53 с. Exit code: 0.")

    add_heading(doc, "9. Скриншоты фактического запуска")
    for filename, caption in SCREENSHOTS:
        add_shot(doc, filename, caption)

    add_heading(doc, "10. Заключение")
    add_para(
        doc,
        "Контрольная точка №2 выполнена: реализованы отдельные автотесты для окон/"
        "вкладок и iframe на Selenium 4 + pytest, применены явные ожидания и "
        "проверяемые assertions, собраны подлинные скриншоты и подтверждена "
        "регрессия КТ 01. Локальные HTML-фикстуры обеспечивают повторяемость при "
        "сохранении связи с публичным демонстрационным сайтом the-internet.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    required = [
        "КТ №2. Тестирование функционала веб-приложения с применением техник тестирования",
        "Цель работы",
        "Требования задания",
        "6 passed in 21.49s",
        "Заключение",
    ]
    missing = [item for item in required if item not in text]
    if missing:
        raise AssertionError(f"DOCX missing sections: {missing}")

    # Count picture elements in the body (identical PNG bytes may share one image part).
    picture_count = 0
    for shape in doc.inline_shapes:
        picture_count += 1
    unique_images = [rel for rel in doc.part.rels.values() if rel.reltype == RT.IMAGE]
    if picture_count < len(SCREENSHOTS):
        raise AssertionError(
            f"Expected >= {len(SCREENSHOTS)} picture elements, got {picture_count}"
        )
    if not unique_images:
        raise AssertionError("No embedded image parts found")
    print(
        f"VERIFY_OK size={path.stat().st_size} "
        f"pictures={picture_count} unique_images={len(unique_images)}"
    )


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    verify(out)
