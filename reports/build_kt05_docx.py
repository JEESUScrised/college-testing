"""Generate reports/KT05.docx from verified multi-run KT05 evidence."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot, add_table, set_run_font
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt05"
OUTPUT = Path(__file__).resolve().parent / "KT05.docx"

SCREENSHOTS = [
    ("01_home_opened.png", "Рисунок 1 — Главная VDNH (актуально после прогона №3)"),
    ("02_news_section_loaded.png", "Рисунок 2 — Раздел /news/"),
    ("04_news_cards.png", "Рисунок 3 — Карточки новостей"),
    ("06_article_opened.png", "Рисунок 4 — Страница статьи"),
    ("08_back_to_news.png", "Рисунок 5 — Возврат в новости"),
    ("09_before_show_more.png", "Рисунок 6 — До «Показать еще»"),
    ("10_after_show_more.png", "Рисунок 7 — После «Показать еще»"),
    ("12_home_to_news_menu.png", "Рисунок 8 — Навигация меню"),
]

# ID, автотест, прогон №1, №2, №3
RESULTS = [
    ("TC-05-01", "test_home_page_opens", "PASS", "—", "PASS"),
    ("TC-05-02", "test_news_section_loads", "PASS", "—", "PASS"),
    ("TC-05-03", "test_news_heading_visible", "PASS", "—", "PASS"),
    ("TC-05-04", "test_news_cards_have_meaningful_content", "PASS", "—", "PASS"),
    ("TC-05-05", "test_news_cards_have_distinct_article_links", "PASS", "—", "PASS"),
    ("TC-05-06", "test_open_article_from_card", "PASS", "—", "PASS"),
    ("TC-05-07", "test_article_heading_and_details", "PASS", "—", "PASS"),
    ("TC-05-08", "test_navigate_back_to_news_section", "PASS", "—", "PASS"),
    ("TC-05-09", "test_show_more_loads_additional_cards", "FAIL", "PASS", "PASS"),
    ("TC-05-10", "test_news_logo_and_menu_navigation", "PASS", "—", "PASS"),
]


def build() -> Path:
    doc = Document()
    setup_a4(doc)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №5. Тестирование функционала сайта и составление тестовой документации",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке", center=True)
    add_para(doc, "Дата обновления отчёта: 2026-10-11", center=True)

    add_heading(doc, "1. Цель и требования")
    add_para(
        doc,
        "В этой КТ мне нужно было прогнать функциональные сценарии по публичному "
        "разделу новостей ВДНХ и оформить нормальную тестовую документацию. "
        "Как ориентир по структуре смотрел репозиторий "
        "https://github.com/sa-teach/selenium-tests — локаторы, понятное дело, "
        "сверял уже с живой вёрсткой сайта.",
    )

    add_heading(doc, "2. Обоснование выбора сайта")
    add_para(doc, "URL: https://vdnh.ru/ и https://vdnh.ru/news/")
    add_para(
        doc,
        "Взял портал ВДНХ, потому что там можно спокойно гонять read-only сценарии "
        "без регистрации. Лента новостей настоящая, есть подгрузка «Показать ещё», "
        "и по сложности хватает, чтобы нормально разложить Page Object. "
        "Нагружать чужой сайт не стал — смысла нет и этически так нельзя. "
        "Подробнее расписал в docs/kt05/website-selection.md.",
    )
    add_para(
        doc,
        "По смыслу сайт информирует посетителей о событиях. Для меня критично было "
        "проверить, что лента открывается, карточки кликаются, статья читается "
        "и подгрузка дополнительных новостей реально работает.",
    )

    add_heading(doc, "3. Среда и инструменты")
    add_para(doc, "Windows; Python 3.13; pytest 8.4.2; Selenium 4.50.0; Chrome headed.")
    add_para(
        doc,
        "Чтобы не зависать на тяжёлой загрузке, поставил pageLoadStrategy=none "
        "и ждал готовности через DOM. Запускал без VPN скриптом "
        "scripts/run_kt05_no_vpn.ps1 — с VPN сайт иногда вёл себя странно.",
    )
    add_para(
        doc,
        "Page Object лежит в selenium/pages/vdnh_pages.py, тесты — "
        "selenium/tests/test_kt05_vdnh_news.py.",
    )

    add_heading(doc, "4. Методика и критерии")
    add_para(
        doc,
        "Black-box — обычные пользовательские сценарии: открыл, кликнул, прочитал. "
        "Gray-box — смотрел DOM, CSS, URL и readyState, когда нужно было понять, "
        "что страница уже догрузилась.",
    )
    add_para(doc, "White-box backend: не применим (закрытый сервер VDNH).")
    add_para(
        doc,
        "По рискам расставил так: P0 — лента и статья, P1 — show_more, "
        "P2 — логотип и меню.",
    )
    add_para(
        doc,
        "Входные условия: Chrome, venv, сайт доступен. Выход: не меньше 8 автотестов, "
        "папка docs/kt05/, DOCX и честные результаты без подгона. "
        "План целиком — в docs/kt05/test-plan.md.",
    )

    add_heading(doc, "5. Трассируемость требований преподавателя")
    add_para(
        doc,
        "1) Выбор сайта — website-selection.md и §2. "
        "2) Подходы black/gray/white — §4. "
        "3) Scope и entry/exit — test-plan.md. "
        "4) Сценарии — test-cases.md плюс таблица ниже.",
    )
    add_para(
        doc,
        "5) Исполнение — §7. 6) Дефекты — §8 и defect-analysis.md "
        "(баг самого сайта в итоге не подтвердился). "
        "7) Рекомендации — §11. 8) Итоговый отчёт — этот DOCX. "
        "Матрица: docs/kt05/requirements-traceability.md.",
    )

    add_heading(doc, "6. Таблица тест-кейсов (три прогона)")
    add_para(
        doc,
        "Важно: прогоны №1 и №2 от 2026-10-09 — это два отдельных запуска. "
        "Их нельзя склеивать в один «красивый» 10/10. "
        "Прогон №3 от 2026-10-11 — уже новый полный suite после правки теста.",
        bold=True,
    )
    add_table(
        doc,
        ["ID", "Автотест", "№1 полный", "№2 show_more", "№3 полный"],
        RESULTS,
    )

    add_heading(doc, "7. Фактические результаты прогонов")
    add_para(
        doc,
        "Прогон №1 (полный, 2026-10-09 примерно в 12:57): "
        "9 PASS / 1 FAIL за 98.63 с, exit code 1.",
    )
    add_code(
        doc,
        "FAILED test_show_more_loads_additional_cards\n"
        "TimeoutException in click_show_more() after 35s wait\n"
        "=================== 1 failed, 9 passed in 98.63s ====================",
    )
    add_para(
        doc,
        "Прогон №2 (только show_more, 2026-10-09 примерно в 13:10): "
        "1 PASS за 11.60 с, exit code 0.",
    )
    add_code(
        doc,
        "test_show_more_loads_additional_cards PASSED\n"
        "============================= 1 passed in 11.60s =============================",
    )
    add_para(
        doc,
        "Прогон №3 (полный после remediation, 2026-10-11 примерно в 01:16): "
        "10 PASS за 50.63 с, exit code 0.",
    )
    add_code(
        doc,
        "collected 10 items\n"
        "============================= 10 passed in 50.63s =============================\n"
        "JUnit: selenium/artifacts/kt05/junit_full_20261011_011634.xml\n"
        "HTML: selenium/artifacts/kt05/report_full_20261011_011634.html",
    )

    add_heading(doc, "8. Анализ FAIL, follow-up и remediation")
    add_para(
        doc,
        "В первом прогоне кнопка «Показать еще» на странице была — я её видел. "
        "Упал уже на ожидании, что контент вырастет после клика. "
        "Старый код смотрел только на cards_count, а после scroll элемент "
        "иногда становился stale. Типичная ловушка Selenium, если честно.",
    )
    add_para(
        doc,
        "Перед вторым прогоном я поправил Page Object: кнопку ищу заново, "
        "а успех считаю, если стало больше карточек, появились новые article links "
        "или сменился URL. Точечный follow-up прошёл, скриншот "
        "10_after_show_more.png это подтверждает.",
    )
    add_para(
        doc,
        "Третий прогон уже полный: 10/10 PASS за 50.63 с. "
        "На баг сайта это не похоже — скорее слабое условие в тесте. "
        "Исторические 9 PASS / 1 FAIL и отдельный follow-up оставляю как есть, "
        "без «склейки» цифр.",
    )

    add_heading(doc, "9. Архитектура Page Object")
    add_para(doc, "VdnhHomePage — главная и переход в новости.")
    add_para(doc, "VdnhNewsListPage — лента, карточки, «Показать еще», логотип.")
    add_para(doc, "VdnhNewsArticlePage — статья и возврат в /news/.")
    add_para(
        doc,
        "Локаторы и действия держу в PO, а assertions оставляю в тестах — "
        "так проще читать и править.",
    )

    add_heading(doc, "10. Скриншоты")
    for name, caption in SCREENSHOTS:
        add_shot(doc, SHOTS / name, caption)

    add_heading(doc, "11. Рекомендации")
    add_para(doc, "По продукту: подтверждённых дефектов в разделе новостей нет.")
    add_para(
        doc,
        "По процессу: не смешивать исторический полный прогон и точечный follow-up; "
        "для подгрузки лучше составные условия успеха; сайту не помешали бы "
        "data-testid на динамических кнопках. Локаторы всё равно приходится "
        "перепроверять, если вёрстка поедет. См. docs/kt05/recommendations.md.",
    )

    add_heading(doc, "12. Заключение")
    add_para(
        doc,
        "По КТ 05 собрал 10 автотестов на https://vdnh.ru/news/ с Page Object "
        "и нормальной документацией (выбор сайта, методика, дефекты, рекомендации). "
        "Сначала полный прогон дал 9 PASS / 1 FAIL, потом show_more отдельно "
        "прошёл (1 passed in 11.60s), а после правки уже весь suite — "
        "10 passed in 50.63s. Баг сайта не подтвердился. "
        "White-box backend: не применим. В целом для учебной работы получилось "
        "достаточно близко к реальному тестированию публичного портала.",
    )

    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    required = [
        "КТ №5. Тестирование функционала сайта и составление тестовой документации",
        "9 PASS / 1 FAIL",
        "1 passed in 11.60s",
        "10 passed in 50.63s",
        "https://vdnh.ru/news/",
        "Обоснование выбора сайта",
        "Рекомендации",
        "Заключение",
        "White-box backend: не применим",
    ]
    missing = [r for r in required if r not in text]
    if missing:
        raise AssertionError(missing)
    if len(doc.tables) < 1:
        raise AssertionError("Expected results table")
    pictures = list(doc.inline_shapes)
    images = [r for r in doc.part.rels.values() if r.reltype == RT.IMAGE]
    if len(pictures) < 6:
        raise AssertionError(f"pictures={len(pictures)}")
    if not images:
        raise AssertionError("no images")
    print(f"VERIFY_OK size={path.stat().st_size} pictures={len(pictures)} unique={len(images)} tables={len(doc.tables)}")


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    verify(out)
