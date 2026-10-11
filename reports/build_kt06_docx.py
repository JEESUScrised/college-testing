"""Generate reports/KT06.docx from verified Appium KT06 evidence (Russian)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot, add_table, set_run_font
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt06"
OUTPUT = Path(__file__).resolve().parent / "KT06.docx"

SCREENSHOTS = [
    ("01_app_home.png", "Рисунок 1 — Главный экран ApiDemos"),
    ("02_app_menu.png", "Рисунок 2 — Меню App"),
    ("03_alert_dialogs.png", "Рисунок 3 — Alert Dialogs"),
    ("04_list_dialog.png", "Рисунок 4 — List dialog"),
    ("05_text_entry_dialog.png", "Рисунок 5 — Text Entry dialog"),
    ("06_back_to_app_menu.png", "Рисунок 6 — Возврат Back на меню App"),
    ("07_back_to_home.png", "Рисунок 7 — Возврат Back на Home"),
    ("08_textfields.png", "Рисунок 8 — Views → TextFields (ввод текста)"),
    ("09_dialog_before_cancel.png", "Рисунок 9 — Text Entry до Cancel"),
    ("10_dialog_after_cancel.png", "Рисунок 10 — Text Entry после Cancel"),
    ("11_textfields_cleared.png", "Рисунок 11 — TextFields после clear"),
    ("12_home_a11y.png", "Рисунок 12 — Home accessibility labels"),
]

# Historical 7/7 preserved; remediation adds TC-08..10
RESULTS_FINAL = [
    ("TC-06-01", "test_app_starts_and_shows_home", "PASS"),
    ("TC-06-02", "test_navigate_to_app_menu", "PASS"),
    ("TC-06-03", "test_open_alert_dialogs_screen", "PASS"),
    ("TC-06-04", "test_list_dialog_interaction", "PASS"),
    ("TC-06-05", "test_text_entry_dialog_input", "PASS"),
    ("TC-06-06", "test_back_returns_to_previous_screen", "PASS"),
    ("TC-06-07", "test_views_textfields_input", "PASS"),
    ("TC-06-08", "test_text_entry_dialog_cancel", "PASS"),
    ("TC-06-09", "test_textfields_empty_after_clear", "PASS"),
    ("TC-06-10", "test_home_menu_accessibility_labels", "PASS"),
]


def build() -> Path:
    doc = Document()
    setup_a4(doc)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №6. Автоматизация мобильного тестирования Android (Appium)",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке", center=True)
    add_para(doc, "Дата обновления: 2026-10-11", center=True)

    add_heading(doc, "1. Цель и выбор приложения")
    add_para(
        doc,
        "Задача КТ — поднять UI-автотесты на официальном sample ApiDemos через "
        "Appium 3 и UiAutomator2. Package: io.appium.android.apis, APK v6.0.18.",
    )
    add_para(
        doc,
        "Почему именно ApiDemos: это публичный sample, его легко поставить заново, "
        "accessibility id стабильные, экраны предсказуемые. Минус тоже понятный — "
        "это не production-приложение, поэтому платежи и бизнес-auth здесь "
        "NOT APPLICABLE. Подробнее — mobile/docs/kt06-test-plan.md.",
    )

    add_heading(doc, "2. Среда выполнения (новый ПК)")
    add_para(doc, "ОС: Windows 10 (19045), AMD64, AMD Ryzen 5 5600, ~16 ГБ RAM.")
    add_para(doc, "Ускорение эмулятора: AEHD 2.2 (usable).")
    add_para(doc, "Python 3.13.2; pytest 8.4.2; Appium-Python-Client 5.3.1.")
    add_para(doc, "Node.js v22.14.0 / npm 10.9.2; локальный Appium 3.8.0; uiautomator2@8.7.0.")
    add_para(doc, "JAVA_HOME: Android Studio JBR 21 (C:\\Program Files\\Android\\Android Studio\\jbr).")
    add_para(doc, "ANDROID_HOME: C:\\Users\\user\\AppData\\Local\\Android\\Sdk.")
    add_para(doc, "AVD: Medium_Phone_API_36.1 (API 36 / Android 16), udid emulator-5554.")
    add_para(
        doc,
        "APK: ApiDemos-debug.apk v6.0.18, SHA-256 "
        "A9EECF37B26CD084855C530DB81C2BB1B91F4C1B095A04F47AA7C20E2791F686.",
    )
    add_para(
        doc,
        "Отдельно: AVD KT06_API33 и system image API 33 я не создавал — "
        "взял уже стоящий эмулятор (вариант B). Физический телефон в этой КТ "
        "не использовал, всё крутил только на эмуляторе.",
    )

    add_heading(doc, "3. Команды запуска")
    add_code(
        doc,
        r""".\scripts\start_kt06_services.ps1 -AvdName Medium_Phone_API_36.1
.\scripts\run_kt06.ps1 -Udid emulator-5554
.\scripts\collect_kt06_artifacts.ps1""",
    )

    add_heading(doc, "4. Хронология прогонов")
    add_para(
        doc,
        "Сначала прогнал один диагностический тест: "
        "test_app_starts_and_shows_home — PASS за 14.94 с. Так убедился, "
        "что Appium вообще видит приложение.",
    )
    add_para(
        doc,
        "Прогон №1 (полный suite до правки scroll): 6 PASS / 1 FAIL за 56.33 с. "
        "Упал test_views_textfields_input — TimeoutException на пункте «TextFields». "
        "На API 36 он просто был ниже видимой области списка Views.",
    )
    add_para(
        doc,
        "Разбирался так: UiScrollable.scrollIntoView(description=\"TextFields\") "
        "элемент показывает. После этого в BaseMobilePage добавил click_menu_item "
        "с прокруткой, если пункта нет на экране.",
    )
    add_para(doc, "Точечный повтор бывшего FAIL: PASS за 15.13 с.")
    add_para(
        doc,
        "Прогон №2 (полный suite после фикса, эмулятор не перезапускал): "
        "7 PASS за 44.63 с, exit 0. Артефакты: "
        "mobile/artifacts/pytest_kt06_20261009_211440.log, "
        "junit_kt06_20261009_211440.xml, report_kt06_20261009_211440.html.",
    )

    add_heading(doc, "5. Результаты прогонов")
    add_para(
        doc,
        "Исторический финал на 2026-10-09: 7 PASS за 44.63 с — эти цифры оставляю. "
        "Потом на remediation (2026-10-11) докинул Cancel, clear/hint и a11y — "
        "получилось 10 PASS за 67.70 с.",
    )
    add_table(doc, ["ID", "Автотест", "Результат"], RESULTS_FINAL)
    add_code(
        doc,
        "============================= 7 passed in 44.63s =============================\n"
        "======================== 10 passed in 67.70s (0:01:07) ========================",
    )

    add_heading(doc, "6. Типы тестирования и платформа")
    add_para(
        doc,
        "Functional: запуск, меню, диалоги, ввод, Back, TextFields — всё прогнал.",
    )
    add_para(
        doc,
        "Negative/dismiss: Cancel в Text Entry; clear в TextFields "
        "(после очистки остаётся hint — так Android и задумано).",
    )
    add_para(doc, "Accessibility: content-desc App/Views на Home.")
    add_para(
        doc,
        "Usability: смотрел по скриншотам и навигации, отдельную UX-сессию "
        "не устраивал.",
    )
    add_para(
        doc,
        "Performance: 67.70 с на 10 кейсов — это ориентир моей среды, "
        "не SLA продукта.",
    )
    add_para(
        doc,
        "Compatibility: только эмулятор Android API 36 "
        "(Medium_Phone_API_36.1). Реальный девайс и другие API — NOT RUN.",
    )
    add_para(
        doc,
        "На живом телефоне DPI и тайминги могут отличаться, поэтому "
        "результаты с эмулятора на физическое устройство без проверки "
        "переносить нельзя.",
    )

    add_heading(doc, "7. Дефекты и рекомендации")
    add_para(
        doc,
        "Подтверждённых багов в ApiDemos нет. INC-06-01 — это был FAIL на TextFields "
        "из-за отсутствия scroll в тесте, поправил в Page Object. "
        "После clear EditText показывает hint («hint text») — ожидаемое поведение "
        "Android, не баг приложения. На будущее: прогнать на real device и другом API, "
        "и не писать в отчёте про payment/auth для demo-приложения.",
    )

    add_heading(doc, "8. Архитектура")
    add_para(doc, "Page Object: mobile/pages/; тесты: mobile/tests/test_kt06_appium.py.")
    add_para(
        doc,
        "Между тестами возвращаю приложение на стартовый экран через "
        "mobile: startActivity → .ApiDemos, чтобы кейсы меньше влияли друг на друга.",
    )

    add_heading(doc, "9. Скриншоты")
    for name, caption in SCREENSHOTS:
        add_shot(doc, SHOTS / name, caption, width_cm=9.0)

    add_heading(doc, "10. Заключение")
    add_para(
        doc,
        "КТ 06 закрыл на эмуляторе Medium_Phone_API_36.1: стратегия, типы тестирования, "
        "кейсы с фактическими результатами, разбор дефекта теста и рекомендации — всё "
        "в документации. Исторически было 7 passed in 44.63s, после remediation — "
        "10 passed in 67.70s. Платежные сценарии NOT APPLICABLE. "
        "Физический телефон не подключал; push в репозиторий не делал. "
        "Самое полезное из этой КТ — понять, что без scrollIntoView на длинных списках "
        "Appium легко ловит «ложный» FAIL.",
    )

    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    required = [
        "КТ №6. Автоматизация мобильного тестирования Android (Appium)",
        "7 PASS",
        "7 passed in 44.63s",
        "10 passed in 67.70s",
        "Medium_Phone_API_36.1",
        "NOT APPLICABLE",
        "Заключение",
    ]
    missing = [r for r in required if r not in text]
    if missing:
        raise AssertionError(missing)
    if len(doc.tables) < 1:
        raise AssertionError("Expected results table")
    pictures = list(doc.inline_shapes)
    images = [r for r in doc.part.rels.values() if r.reltype == RT.IMAGE]
    if len(pictures) < 10:
        raise AssertionError(f"pictures={len(pictures)}")
    # Word may dedupe identical PNG bytes into fewer image parts
    if len(images) < 8:
        raise AssertionError(f"images={len(images)}")
    print(
        f"VERIFY_OK size={path.stat().st_size} pictures={len(pictures)} "
        f"unique={len(images)} tables={len(doc.tables)}"
    )


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    verify(out)
