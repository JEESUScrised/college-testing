"""Generate reports/KT08.docx — Pillow history + public site + Allure compliance."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from docx_helpers import setup_a4, add_heading, add_para, add_code, add_shot, add_table, set_run_font
from report_title import add_title_fields

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots" / "kt08"
BASELINES = ROOT / "selenium" / "baselines" / "kt08"
ARTIFACTS = ROOT / "selenium" / "artifacts" / "kt08"
OUTPUT = Path(__file__).resolve().parent / "KT08.docx"

PILLOW_RESULTS = [
    ("TC-08-01", "match baseline", "0.0000%", "PASS"),
    ("TC-08-02", "text/color change", "1.1082%", "PASS"),
    ("TC-08-03", "layout shift", "5.0647%", "PASS"),
    ("TC-08-04", "significant diff", "99.8457%", "PASS"),
    ("TC-08-05", "dimension mismatch", "n/a", "PASS"),
]

PUBLIC_RESULTS = [
    ("TC-08-P01", "Login page / branding / form", "PASS"),
    ("TC-08-P02", "Inventory after login", "PASS"),
    ("TC-08-P03", "Sort dropdown", "PASS"),
    ("TC-08-P04", "Add to cart / badge", "PASS"),
]

COMPLIANCE = [
    ("1. Выбрать публичное веб-приложение и UI-элементы", "Sauce Demo (saucedemo.com); login, inventory, dropdown, cart", "docs + public_*.png", "PASS"),
    ("2. Настроить Selenium + test framework", "pytest + Selenium 4.50 Chrome viewport 1400×1000", "conftest / run_kt08.ps1", "PASS"),
    ("3. Интегрировать Allure", "allure-pytest 2.16.2 + CLI 2.37.0", "allure-results / allure-report", "PASS"),
    ("4. Автотесты скриншотов UI/страниц", "4 public + 5 Pillow tests", "test_kt08_*.py", "PASS"),
    ("5. Выполнить тесты и собрать скриншоты", "Full 9 passed in 34.64s", "screenshots/kt08/", "PASS"),
    ("6. Сгенерировать Allure HTML со скриншотами", "allure generate → index.html; PNG в шагах", "allure-report + allure_*.png", "PASS"),
    ("7. Анализ UI + отчёт с выводами", "Этот DOCX + docs/kt08/", "reports/KT08.docx", "PASS"),
]


def build() -> Path:
    doc = Document()
    setup_a4(doc)

    add_title_fields(doc, add_para)
    add_para(doc, "")
    add_para(
        doc,
        "КТ №8. Скриншотное тестирование веб-приложений и Allure",
        bold=True,
        center=True,
    )
    add_para(doc, "Отчёт по контрольной точке (compliance remediation)", center=True)
    add_para(doc, "Дата: 2026-10-10", center=True)

    add_heading(doc, "1. Цель задания (из презентации преподавателя)")
    add_para(
        doc,
        "По презентации нужно разобраться со скриншотным тестированием веб-приложений "
        "и связать результаты с Allure: выбрать публичный сайт, настроить Selenium, "
        "подключить Allure, прогнать скриншотные автотесты, собрать HTML-отчёт "
        "и написать свой разбор.",
    )

    add_heading(doc, "2. Публичный сайт и обоснование")
    add_para(
        doc,
        "Выбрал https://www.saucedemo.com/ — публичное demo Swag Labs. "
        "Там есть login, каталог, dropdown сортировки и корзина: как раз то, "
        "что удобно снимать скриншотами и вешать в Allure.",
    )
    add_para(
        doc,
        "Сначала смотрел https://the-internet.herokuapp.com/, но под Chrome automation "
        "на моей машине он не открылся нормально: после get() DOM оставался пустым "
        "(challenge/bot script). В отчёт это записал отдельно. Localhost вместо "
        "публичного сайта не подставлял — требование презентации так не закрыть.",
        bold=True,
    )

    add_heading(doc, "3. Сценарии Selenium на публичном сайте")
    add_table(doc, ["ID", "Сценарий", "Статус"], PUBLIC_RESULTS)
    add_para(
        doc,
        "Файл: selenium/tests/test_kt08_public_allure.py. Viewport зафиксировал "
        "1400×1000, pageLoadStrategy=eager, ждал через явные ожидания, "
        "без time.sleep.",
    )

    add_heading(doc, "4. Архитектура Allure")
    add_para(
        doc,
        "allure-pytest 2.16.2 пишет сырые результаты в "
        "selenium/artifacts/kt08/allure-results/. "
        "Allure CLI 2.37.0 на Java 21 собирает HTML: allure generate … -o allure-report. "
        "В тестах использую @allure.title/description/feature/step и "
        "allure.attach.file для PNG — так скриншот виден прямо в шаге.",
    )
    add_code(
        doc,
        r""".\scripts\run_kt08.ps1 -Mode Full
# → allure-results + allure-report\index.html""",
    )

    add_heading(doc, "5. Скриншоты публичного сайта")
    for name, cap in [
        ("public_01_login_page.png", "Рис. P1 — Login page"),
        ("public_02_inventory.png", "Рис. P2 — Inventory / Products"),
        ("public_03_sort_dropdown.png", "Рис. P3 — Sort dropdown (Z to A)"),
        ("public_04_add_to_cart.png", "Рис. P4 — Add to cart / badge=1"),
    ]:
        add_shot(doc, SHOTS / name, cap)

    add_heading(doc, "6. Allure attachments и HTML-отчёт")
    add_para(
        doc,
        "В allure-results лежит 9 test result JSON; PNG и text attachments "
        "сидят во вложенных steps — это я отдельно проверял парсером. "
        "HTML-отчёт открывается, в title написано «Allure Report».",
    )
    add_shot(doc, SHOTS / "allure_01_overview.png", "Рис. A1 — Allure overview")
    add_shot(doc, SHOTS / "allure_02_suites.png", "Рис. A2 — Allure suites/navigation")

    add_heading(doc, "7. Результаты выполнения")
    add_para(doc, "DiagnosticPublic: 1 passed in 3.84s.")
    add_para(doc, "Public only: 4 passed in 21.27s.")
    add_code(doc, "============================= 9 passed in 34.64s ==============================")
    add_para(
        doc,
        "Full suite (Pillow + Public) от 2026-10-10. "
        "Лог: pytest_kt08_20261010_144813.log.",
    )

    add_heading(doc, "8. Анализ UI Sauce Demo")
    add_para(
        doc,
        "На login page сразу видны поля user-name/password и подсказка "
        "accepted usernames — для учёбы удобно, не надо искать тестовые учётки. "
        "В inventory стабильные data-test атрибуты, сортировка через обычный "
        "<select>, badge корзины обновляется после Add to cart. "
        "Нюанс: после select React перерисовывает DOM, элемент надо искать заново. "
        "the-internet.herokuapp.com у меня под automation не зашёл — оставляю "
        "как ограничение среды.",
    )
    add_para(
        doc,
        "Из практики: для скриншотных тестов лучше брать сайты без bot-challenge, "
        "фиксировать viewport и крепить PNG в шаги Allure, а не только класть "
        "файлы в папку screenshots.",
    )

    add_heading(doc, "9. Сохранённый контур Pillow (исторический)")
    add_para(
        doc,
        "Параллельно оставил локальные HTML-фикстуры и compare.py на Pillow. "
        "Исторический прогон: 5 passed in 13.71s (2026-10-09). "
        "В remediation эти пять кейсов тоже вошли в Full suite и снова прошли; "
        "diff/overlay дополнительно аттачатся в Allure.",
    )
    add_table(doc, ["ID", "Сценарий", "Diff%", "Статус"], PILLOW_RESULTS)
    add_shot(doc, BASELINES / "home_baseline.png", "Рис. L1 — Approved baseline", width_cm=9.0)
    add_shot(doc, SHOTS / "01_home_actual.png", "Рис. L2 — Actual match", width_cm=9.0)
    if (ARTIFACTS / "02_text_color_overlay.png").is_file():
        add_shot(doc, ARTIFACTS / "02_text_color_overlay.png", "Рис. L3 — Overlay text/color", width_cm=9.0)

    add_heading(doc, "10. Compliance checklist (7 требований)")
    add_table(doc, ["Требование", "Реализация", "Доказательство", "Статус"], COMPLIANCE)

    add_heading(doc, "11. Заключение")
    add_para(
        doc,
        "КТ 08 закрыл под требования презентации: публичный saucedemo.com, "
        "Allure HTML со скриншотами в шагах, полный прогон 9 passed in 34.64s. "
        "Локальный Pillow-контур (5 passed in 13.71s) оставил как детерминированную "
        "проверку сравнения картинок — на публичном сайте верстка может чуть "
        "плавать, а локальные фикстуры для diff удобнее. "
        "Если коротко увиденное: Allure полезен именно когда PNG висят в steps, "
        "иначе отчёт превращается в сухой список PASS без контекста.",
    )

    doc.save(OUTPUT)
    return OUTPUT


def verify(path: Path) -> None:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    required = [
        "КТ №8. Скриншотное тестирование веб-приложений и Allure",
        "9 passed in 34.64s",
        "saucedemo.com",
        "Compliance checklist",
        "5 passed in 13.71s",
    ]
    missing = [r for r in required if r not in text]
    if missing:
        raise SystemExit(f"DOCX missing: {missing}")
    n_pics = sum(1 for r in doc.part.rels.values() if "image" in r.reltype)
    print(f"OK {path} size={path.stat().st_size} images={n_pics} chars={len(text)}")


if __name__ == "__main__":
    out = build()
    verify(out)
